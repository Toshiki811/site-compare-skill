#!/usr/bin/env python3
"""sc_book.py — 拠点比較の表（out/拠点比較.xlsx）を書き出す（工程 8。機械だけ。LLM は書かない・直さない）。

  python3 scripts/sc_book.py -w <比較の作業フォルダ>                    # 書き出す（確認事項一覧の回答・回答者・回答日は、いまのブックから取り込んで戻す）
  python3 scripts/sc_book.py -w <比較の作業フォルダ> --template <様式.xlsx>   # お客様の様式で
  python3 scripts/sc_book.py default                                   # 既定のテンプレート（assets/拠点比較_テンプレート.xlsx）を作り直す

シート:
  概要                 拠点・基準・行とタスクの数・区分ごとの数・照合の支持率・複雑性の決め方と境
  拠点タスク一覧       ① 1 行＝1 タスク。業務分類・業務種別・タスク・区分（共通／差異あり／拠点固有）と、拠点ごとの列（記号＋その拠点の作業の名前。
                       内訳は改行で）・差の種類・差の内容・粒度・片方だけ記載・確認事項
  拠点別の明細         1 行＝タスク × 拠点。作業の名前・担当・システム・帳票・出典・推定・粒度・差
  業務種別の対応       1 行＝業務種別。拠点ごとに、当てたフローと当て方
  複雑性ヒートマップ   ③ 1 行＝業務種別。拠点ごとに、大中小（塗り）・分岐と受け渡しの数・点数（判定に使った粒度と、もう一方の粒度の数）・差の大きさ
  複雑性の明細         数えた分岐・受け渡しを 1 件ずつ（どちらの粒度で数えたか）
  追加の資料の依頼     拠点 × 業務種別ごとに、粗く書いた・記載なしのタスクと、活かせていない作業の数、依頼する資料（比べられるようになるタスクの多い順）
  拠点間の用語対応     比べる名前と、拠点ごとの呼び方・まとめた根拠
  確認事項一覧         比較の確認事項（Q 番号は引き継ぐ）。回答・回答者・回答日は、人が埋める（書き出し直しても残る）
  読み方               記号・色・数え方の説明
拠点の列は、テンプレートの「①」の列を、拠点の数だけ機械が繰り返す（見出しは拠点の名前。基準が左）。
"""
import argparse
import datetime
import json
import pathlib
import sys

from sc_common import ASSETS, BOOK, A, digest, load, result_path, save, sites_of, try_load, work_dir
from sc_rows import Model

import table_template as TT                                          # asis の表の部品

TEMPLATE = "拠点比較_テンプレート.xlsx"
T_TASK, T_DETAIL, T_FLOW, T_HEAT, T_HEATD, T_REQ, T_TERMS, T_Q = ("拠点タスク一覧", "拠点別の明細", "業務種別の対応", "複雑性ヒートマップ", "複雑性の明細",
                                                                  "追加の資料の依頼", "拠点間の用語対応", "確認事項一覧")
HEADS = {
    T_TASK: ["No", "業務分類", "業務種別", "タスク", "区分", "差の種類", "差の内容", "粒度", "片方だけ記載", "確認事項"],
    T_DETAIL: ["No", "業務種別", "タスク", "拠点", "記号", "その拠点の作業", "担当", "システム", "帳票", "出典", "推定", "粒度", "差"],
    T_FLOW: ["業務分類", "業務種別", "タスクの数", "共通", "差異あり", "拠点固有"],
    T_HEAT: ["業務分類", "業務種別", "差の大きさ"],
    T_HEATD: ["業務種別", "拠点", "数え方", "種類", "タスク", "中身"],
    T_REQ: ["拠点", "業務種別", "比べられるようになるタスクの数", "粗く書いたタスク", "記載なしのタスク", "活かせていない作業", "判断", "受け渡し", "システム", "帳票", "依頼する資料", "説明"],
    T_TERMS: ["種類", "比べる名前", "まとめた根拠", "理由"],
    T_Q: ["番号", "種別", "答えで変わる所", "業務種別", "タスク", "拠点", "質問", "仮置きした内容", "聞く相手"],
}
REPEAT = {T_TASK: "拠点", T_FLOW: "拠点", T_HEAT: "拠点", T_TERMS: "拠点"}
FILLS = [("共通", "common", "E2EFDA", "区分が共通のタスク。複雑性の「小」"), ("差異あり", "diff", "FCE4D6", "区分が差異ありのタスク"),
         ("拠点固有", "unique", "F8CBAD", "区分が拠点固有のタスク"), ("記載なし", "none", "F2F2F2", "その拠点の資料に記載なし・対象外のセル"),
         ("差のある拠点", "mark", "FBE5D6", "その拠点が差に関わるセル（△）"), ("小", "small", "E2EFDA", "複雑性が小"), ("中", "mid", "FFE699", "複雑性が中"),
         ("大", "large", "F8CBAD", "複雑性が大"), ("推定", "inferred", "FCE4D6", "推定の箱を含むセル")]
FONTS = [("記号", "mark", dict(sz=11, b=True), "○ △ — の記号"), ("薄い字", "pale", dict(sz=9, color="808080"), "内訳・粒度の注記"),
         ("出典", "src", dict(sz=8, color="7F7F7F"), "出典の資料とページ"), ("太字", "bold", dict(sz=10, b=True), "タスクの名前・大中小"),
         ("差", "diff", dict(sz=9, color="C55A11"), "差の種類")]
LEGEND = [
    ("このブック", "拠点ごとの as-is の業務フロー（asis-flow-skill で作り、原典照合を済ませたもの）を、拠点をまたいで比べた一覧。図（拠点比較フロー.drawio）と同じデータから機械で作った。"
                   "直すときは、もとのデータ（用語の対応・比べる行・タスクの対応）を直して書き出し直す（このブックに書き込んだ内容は、確認事項一覧の回答の欄を除いて、書き出し直すと消える）。"),
    ("業務分類・業務種別・タスク", "業務分類 › 業務種別（比べる行。拠点のフロー 1 本か、その一部）› タスク（拠点をまたいで同じ作業とした単位）。"
                                   "タスクは、いちばん粗く書いた拠点に合わせる（ほかの拠点が 1 つの作業で書いた所は、詳しい拠点の作業をまとめて 1 つのタスクにし、内訳として残す）。"),
    ("区分", "共通＝比べた拠点のすべてに記載があり、差が無い／差異あり＝2 拠点以上に記載があり、担当・システム・帳票・判断・順序・内容のどれかが違う、または記載の無い拠点がある／"
             "拠点固有＝記載のある拠点が 1 つだけ。"),
    ("○ △ —", "○＝記載あり・差なし／△＝記載あり・差がある（基準と違う。基準が無ければ、記載のある最初の拠点と違う）／—＝資料に記載なし（業務が無いとは言っていない）／対象外＝その拠点で対象外にした業務。"),
    ("差の種類", "担当（レーン）・システム・帳票・判断（行き先が別のタスクに分かれる判断の数）・順序（線でたどれる前後が逆）・内容（やり方・値・対象。根拠つきで書いたもの）・記載なし。"
                 "片方の拠点にしか書いていない担当・システム・帳票は「片方だけ記載」の列に出し、差には数えない。"),
    ("粒度", "「粗」＝その拠点が、ほかの拠点が詳しく書いたタスクを 1 つの作業で書いている（比べるときは粗い側に合わせた）。詳しい拠点の作業は、拠点の列に内訳として残している。"),
    ("複雑性", "分岐＝判断の行き先の数−1 の合計。受け渡し＝担当（レーン）が変わる線の数（システムの自動処理を挟む線は、前後の担当で見る）。"
               "点数＝分岐×重み＋受け渡し×重み。大中小の境と重み・判定に使う粒度は「概要」。そろえた粒度＝タスクの境をまたぐ分岐・受け渡しだけを数える（書き方の細かさに引っ張られない）。"
               "原典の粒度＝拠点のフローのまま数える。どちらの数も出す。"),
    ("差の大きさ", "基準の拠点（基準の無い行は、記載のある拠点の中で点数のいちばん小さい拠点）との点数の差から、大中小。"),
    ("追加の資料の依頼", "粗く書いた拠点・記載なしの拠点ごとに、どの資料があれば、使えていない詳しい情報（ほかの拠点の作業・判断・受け渡し）と比べられるかを、比べられるようになるタスクの多い順に並べた。"),
    ("出典", "資料の題名と、資料に刷られたページ番号。原文の引用は、拠点ごとの as-is の出典対応表にある。"),
]
OVERVIEW = [("比べた拠点", "{拠点}（表と図では、基準の拠点を左に置く）"), ("基準の拠点", "基準: {基準}"), ("比べる行（業務種別）", "{行の数} 行（業務分類 {分類の数}）"),
            ("タスク", "{タスクの数}（共通 {共通}・差異あり {差異あり}・拠点固有 {拠点固有}）"), ("対応の照合の支持率", "支持率 {支持率}"),
            ("複雑性の大中小", "大中小は、{決め方}"), ("大中小の境", "境: {境}"), ("差の大きさの境", "{差の境}（拠点ごとの大中小とは別の列）"),
            ("拠点ごとの as-is", "拠点の写し: {拠点の状態}"),
            ("確認事項", "{確認事項の数} 件（比較のもの。拠点ごとの as-is の確認事項は、拠点の確認事項一覧）"), ("作った日", "作った日: {日付}")]
SPEC = {"sheets": HEADS, "repeat": REPEAT, "values": ["拠点", "基準", "行の数", "分類の数", "タスクの数", "共通", "差異あり", "拠点固有", "支持率", "決め方", "境",
                                                     "差の境", "拠点の状態", "確認事項の数", "日付"],
        "required": {T_TASK: ["No", "タスク", "区分", "拠点"], T_Q: ["番号", "質問"]}, "keys": {T_Q: ["番号"], T_TASK: ["No"]},
        "fills": {label: key for label, key, _c, _w in FILLS}, "fonts": {label: key for label, key, _f, _w in FONTS}}
LOOK = {"font": "游ゴシック", "size": 9, "border": True, "head_fill": "D9E1F2", "head_font": dict(sz=10, b=True), "step_width": 34,
        "order": ["概要", T_TASK, T_HEAT, T_REQ, T_DETAIL, T_FLOW, T_HEATD, T_TERMS, T_Q, "読み方"],
        "manual": {T_Q: ["回答", "回答者", "回答日"]},
        "widths": {T_TASK: [8, 14, 16, 30, 9, 16, 40, 18, 30, 10], T_DETAIL: [8, 16, 28, 14, 6, 40, 16, 18, 16, 30, 6, 14, 40],
                   T_FLOW: [14, 20, 8, 8, 8, 8], T_HEAT: [14, 22, 10], T_HEATD: [18, 14, 12, 10, 26, 60],
                   T_REQ: [14, 18, 10, 18, 18, 10, 8, 8, 8, 8, 48, 60], T_TERMS: [10, 18, 12, 50], T_Q: [8, 10, 12, 16, 22, 14, 60, 30, 16, 40, 12, 12],
                   "概要": [24, 100], "読み方": [24, 110]},
        "freeze": {T_TASK: "E2", T_HEAT: "C2", T_DETAIL: "D2"}, "filter": [T_TASK, T_DETAIL, T_HEATD, T_REQ, T_Q],
        "page": {"orientation": "landscape", "paper": 8, "title_rows": "1:1"},
        "fills": {label: (color, where) for label, _k, color, where in FILLS}, "fonts": {label: (key, font, where) for label, key, font, where in FONTS},
        "static": {"概要": {"head": ["項目", "内容"], "rows": OVERVIEW, "fonts": ["太字", None]},
                   "読み方": {"head": ["項目", "説明"], "rows": LEGEND, "fonts": ["太字", None]}}}
LAYOUT = {T_TASK: ["No", "業務分類", "業務種別", "タスク", "区分", "*", "差の種類", "差の内容", "粒度", "片方だけ記載", "確認事項"],
          T_FLOW: ["業務分類", "業務種別", "タスクの数", "共通", "差異あり", "拠点固有", "*"], T_HEAT: ["業務分類", "業務種別", "*", "差の大きさ"],
          T_TERMS: ["種類", "比べる名前", "*", "まとめた根拠", "理由"]}
SITE_WIDTH = {T_TASK: 34, T_FLOW: 30, T_HEAT: 30, T_TERMS: 18}


def make_template(path):
    """既定のテンプレート: asis の make_default で作り、拠点の列（①）を、決めた位置へ移す"""
    import openpyxl
    from openpyxl.utils import get_column_letter
    TT.make_default(SPEC, LOOK, path)
    wb = openpyxl.load_workbook(path)
    for title, order in LAYOUT.items():
        ws = wb[title]
        cells = {}
        for c in range(1, ws.max_column + 1):
            mark = ws.cell(2, c).value
            key = "*" if mark == "{拠点}" else (mark or "").strip("{}")
            cells[key] = [(ws.cell(r, c).value, ws.cell(r, c)._style, ws.column_dimensions[get_column_letter(c)].width) for r in (1, 2)]
        for c in range(1, ws.max_column + 1):
            for r in (1, 2):
                ws.cell(r, c).value = None
        for c, key in enumerate(order, 1):
            for r, (v, st, w) in zip((1, 2), cells[key]):
                x = ws.cell(r, c)
                x.value, x._style = v, st
            ws.column_dimensions[get_column_letter(c)].width = SITE_WIDTH[title] if key == "*" else cells[key][0][2]
        if ws.auto_filter.ref:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(order))}2"
    wb.save(path)


C = TT.Cell


def site_cell(x, sid, mdl):
    """拠点タスク一覧の、拠点のセル: 記号＋その拠点の作業の名前（内訳は改行）"""
    mk = x["marks"].get(sid, "—")
    mem = x["members"].get(sid) or []
    runs = [(mk, "mark")]
    if mem:
        if len(mem) == 1:
            runs.append((" " + mem[0]["label"], "text"))
        else:
            runs.append((f" 内訳 {len(mem)}", "pale"))
            for m in mem:
                runs.append(("\n・" + m["label"], "text"))
        if sid in x["grain"]["coarse"]:
            runs.append(("\n（粗: ほかの拠点の " + str(max(x["grain"]["counts"].values())) + " 作業に当たる）", "pale"))
        if any(m["inferred"] for m in mem):
            runs.append(("\n（推定を含む）", "pale"))
    elif mk == "—":
        runs.append((" 記載なし", "pale"))
    fill = {"—": "none", "対象外": "none", "△": "mark"}.get(mk)
    return C(runs, fill=fill)


def questions(cmp, mdl, write=True):
    """比較の確認事項（出どころの順）。通し番号は numbers.json の q で引き継ぐ（write=False は、番号を控えない。点検が使う）"""
    out = []
    for origin, path in [("applies", cmp / "applies.json"), ("terms", cmp / "terms.json"), ("axis", cmp / "axis.json")] + \
            [(f"rows/{r['id']}", cmp / "rows" / f"{r['id']}.match.json") for r in mdl.rows()]:
        d = try_load(path)[0]
        for q in (d or {}).get("questions") or []:
            out.append((origin, q))
    nb = load(cmp / "numbers.json", None) or {"next": 1, "assigned": {}}
    nb.setdefault("qnext", 1)
    nb.setdefault("qassigned", {})
    rows = []
    for origin, q in out:
        key = f"{origin}|{q['id']}"
        if key not in nb["qassigned"]:
            nb["qassigned"][key] = f"Q-{nb['qnext']:03d}"
            nb["qnext"] += 1
        rows.append((nb["qassigned"][key], origin, q))
    if write:
        save(cmp / "numbers.json", nb)
    return rows


def book_mark(results, qs):
    """表の中身の印（計算の結果と確認事項）。書き出したときの印と違えば、書き出し直す"""
    return digest([results, [[n, o, q] for n, o, q in qs]], 16)


def sheets(cmp, mdl, results, qs):
    names = [mdl.names[s] for s in mdl.order]
    heads = lambda t: [C(h) for h in HEADS[t]] + ([C(("【基準】" if s == mdl.baseline else "") + mdl.names[s]) for s in mdl.order] if t in REPEAT else [])
    out = {t: [heads(t)] for t in HEADS}
    qof = {}
    for no, origin, q in qs:
        if origin.startswith("rows/"):
            qof.setdefault((origin[5:], q.get("task")), []).append(no)
    tid_name = {}
    for res in results:
        rid = res["row"]
        for x in res["tasks"]:
            tid_name[(rid, x["id"])] = (x["no"], x["name"])
            kinds = [d["kind"] for d in x["diffs"]]
            row = [C(x["no"]), C(res.get("group_name") or "—"), C(res["name"]), C(x["name"], font="bold"), C(x["class"], fill={"共通": "common", "差異あり": "diff", "拠点固有": "unique"}[x["class"]]),
                   C("・".join(dict.fromkeys(kinds)) or "—", font="diff" if kinds else None),
                   C("\n".join(f"{d['kind']}: {d['text']}" for d in x["diffs"] if d["kind"] != "記載なし") or "—"),
                   C(("粗: " + "・".join(mdl.names[s] for s in x["grain"]["coarse"]) + "／詳: " + "・".join(f"{mdl.names[s]} {x['grain']['counts'][s]}" for s in x["grain"]["fine"])) if x["grain"]["coarse"] else "—"),
                   C("\n".join(f"{d['kind']}: {d['text']}" for d in x["onesided"]) or "—"),
                   C("・".join(qof.get((rid, x["id"]), [])) or "—")]
            out[T_TASK].append(row + [site_cell(x, s, mdl) for s in mdl.order])
            for s in mdl.order:
                mem = x["members"].get(s) or []
                mk = x["marks"].get(s, "—")
                out[T_DETAIL].append([C(x["no"]), C(res["name"]), C(x["name"]), C(mdl.names[s]), C(mk, font="mark", fill={"—": "none", "対象外": "none", "△": "mark"}.get(mk)),
                                      C("\n".join(m["label"] for m in mem) or "（記載なし）"), C("\n".join(dict.fromkeys(m["role"] for m in mem)) or "—"),
                                      C("・".join(dict.fromkeys(y for m in mem for y in m["systems"])) or "—"), C("・".join(dict.fromkeys(f for m in mem for f in m["files"])) or "—"),
                                      C([(y, "src") for y in dict.fromkeys(z for m in mem for z in m["src"])] or "—"), C("推定" if any(m["inferred"] for m in mem) else ""),
                                      C("粗" if s in x["grain"]["coarse"] else "詳" if s in x["grain"]["fine"] and x["grain"]["coarse"] else ""),
                                      C("\n".join(f"{d['kind']}: {d['text']}" for d in x["diffs"] if s in d["sites"]) or "—")])
        cls = [x["class"] for x in res["tasks"]]
        out[T_FLOW].append([C(res.get("group_name") or "—"), C(res["name"]), C(str(len(cls))), C(str(cls.count("共通"))), C(str(cls.count("差異あり"))), C(str(cls.count("拠点固有")))]
                           + [C(flow_cell(mdl, rid, s, res)) for s in mdl.order])
        heat = []
        other = "raw" if res["basis"] == "aligned" else "aligned"
        for s in mdl.order:
            e = res["sites"].get(s) or {}
            c = e.get("complexity")
            if not c:
                heat.append(C([("対象外" if e.get("relation") == "out_of_scope" else "記載なし", "pale")], fill="none"))
                continue
            v, o = c[res["basis"]], c[other]
            runs = [(v["level"], "bold"), (f"　分岐 {v['br']}・受け渡し {v['hand']}（点 {v['score']}）", "text"),
                    (f"\n{'原典' if other == 'raw' else 'そろえた'}の粒度: 分岐 {o['br']}・受け渡し {o['hand']}（点 {o['score']}）", "pale")]
            if e.get("grain"):
                runs.append(("\n粒度差（粗く書いた側）", "pale"))
            if e.get("spread"):
                runs.append((f"\n基準との差 {e['spread']['diff']:+d}・差のあるタスク {int(e['spread']['ratio'] * 100)}%", "pale"))
            heat.append(C(runs, fill={"小": "small", "中": "mid", "大": "large"}[v["level"]]))
            for it in c["counted"]:
                out[T_HEATD].append([C(res["name"]), C(mdl.names[s]), C("そろえた粒度" if it["basis"] == "aligned" else "原典の粒度"), C(it["kind"]),
                                     C(it["where"]), C(it["text"])])
        out[T_HEAT].append([C(res.get("group_name") or "—"), C(res["name"]), C(res.get("spread") or "—")] + heat)
        for q in res["requests"]:
            f = q["folded"]
            out[T_REQ].append((q["gain"], mdl.order.index(q["site"]), [C(mdl.names[q["site"]]), C(res["name"]), C(str(q["gain"])), C("・".join(q["coarse"]) or "—"),
                                                                        C("・".join(q["absent"]) or "—"), C(str(f.get("nodes", 0))), C(str(f.get("decisions", 0))),
                                                                        C(str(f.get("handoffs", 0))), C(str(f.get("systems", 0))), C(str(f.get("files", 0))),
                                                                        C("\n".join(q["kinds"])), C(q["text"])]))
    out[T_REQ] = out[T_REQ][:1] + [r for _g, _o, r in sorted(out[T_REQ][1:], key=lambda x: (-x[0], x[1]))]
    for g in mdl.terms.get("groups") or []:
        per = {s: [m["term"] for m in g["members"] if m["site"] == s] for s in mdl.order}
        out[T_TERMS].append([C(dict(actor="担当", system="システム", document="帳票")[g["kind"]]), C(g["name"], font="bold"),
                             C(dict(same_name="名前が同じ", translation="訳", equivalent="言い換え")[g["basis"]]), C(g.get("reason") or "—")]
                            + [C("・".join(per[s]) or "—", fill=None if per[s] else "none") for s in mdl.order])
    for no, origin, q in qs:
        task = tid_name.get((origin[5:], q.get("task"))) if origin.startswith("rows/") else None
        row = origin[5:] if origin.startswith("rows/") else q.get("row")
        out[T_Q].append([C(no), C(q["category"]), C(q["effect"]), C(mdl.row_of.get(row, {}).get("name", "—") if row else "—"),
                         C(f"{task[0]} {task[1]}" if task else "—"), C(mdl.names.get(q.get("site"), "—") if q.get("site") else "—"),
                         C(q["question"]), C(q.get("assumption") or ""), C(q.get("ask") or "")])
    return out


def flow_cell(mdl, rid, sid, res):
    e = res["sites"].get(sid) or {}
    sp = mdl.site_part(rid, sid) or {}
    rel = {"same": "", "part": "（この業務の一部）", "split": "（ほかの業務と 1 本のフローの一部）", "none": "記載なし", "out_of_scope": "対象外"}.get(e.get("relation"), "")
    txt = "\n".join(e.get("flows") or []) + rel
    if sp.get("reason") and e.get("relation") != "same":
        txt += f"\n理由: {sp['reason']}"
    return txt or "—"


def values(cmp, mdl, results, qs):
    import sc_compute
    cfg = sc_compute.config()
    cls = [x["class"] for r in results for x in r["tasks"]]
    rate = "—"
    try:
        import sc_verify
        tot = ok = 0
        for r in results:
            cur = sc_verify.claims(mdl, r["row"])
            rv = try_load(cmp / "review" / f"{r['row']}.review.json")[0] or {}
            got = {v["claim"]: v for v in rv.get("verdicts") or []}
            shown = {v["id"]: v for v in sc_verify.versions(cmp)["versions"]}
            for c, (mark, _b) in cur.items():
                tot += 1
                v = got.get(c)
                ok += bool(v and v["verdict"] == "支持" and shown.get(v["sheet"], {}).get("claims", {}).get(c) == mark)
        rate = f"{ok / tot:.0%}（支持 {ok} ÷ 主張 {tot}）" if tot else "—"
    except Exception:                                                # 照合の記録が崩れていても、表は出す（点検が止める）
        pass
    w = cfg["weights"]
    border = "・".join(f"{lv['name']}: 点 {lv['max']} 以下" if lv["max"] is not None else f"{lv['name']}: それより上" for lv in cfg["levels"])
    state = "・".join(f"{mdl.names[s]}（写し {(mdl.snaps.get(s) or {}).get('made', '—')[:10]}・asis の確認事項 {(mdl.snaps.get(s) or {}).get('questions', 0)} 件）" for s in mdl.order)
    return {"拠点": "・".join(mdl.names[s] for s in mdl.order), "基準": mdl.names.get(mdl.baseline, "なし（拠点を横に並べて、同じ業務をそろえた）"),
            "行の数": str(len(results)), "分類の数": str(len(mdl.axis.get("groups") or [])), "タスクの数": str(len(cls)),
            "共通": str(cls.count("共通")), "差異あり": str(cls.count("差異あり")), "拠点固有": str(cls.count("拠点固有")), "支持率": rate,
            "決め方": f"{'そろえた粒度' if cfg['basis'] == 'aligned' else '原典の粒度'}の、分岐 × {w['branch']} ＋ 受け渡し × {w['handoff']} の点数",
            "境": border, "差の境": f"基準との点数の差が {cfg['spread']['score'][0]} 以上で中・{cfg['spread']['score'][1]} 以上で大",
            "拠点の状態": state, "確認事項の数": str(len(qs)), "日付": datetime.date.today().isoformat()}


def write(cmp, template=None):
    mdl = Model(cmp, sites_of(cmp))
    results = [r for r in (try_load(result_path(cmp, x["id"]))[0] for x in mdl.rows()) if isinstance(r, dict)]
    if not results:
        sys.exit("計算の結果が無い（python3 scripts/sc_compute.py を先に）")
    qs = questions(cmp, mdl)
    book = cmp / "out" / BOOK
    tpl_path = pathlib.Path(template) if template else ASSETS / TEMPLATE
    ans = A.Answers(cmp)
    said = []
    if book.exists():
        try:
            got, notes = TT.manual_back(book, SPEC, tpl_path)
            now = {T_Q: {no: q["question"] for no, _o, q in qs}}
            n = ans.take(BOOK, {k: v for k, v in got.items() if k == T_Q}, now)
            said += notes + ([f"確認事項一覧の回答を取り込んだ: {n} 行（answers.json）"] if n else [])
        except TT.TemplateError as e:
            sys.exit(f"いまのブックが読めない（{e}）。人が入れた回答を消さないよう止める（開いていれば閉じる）")
    ans.save()
    data = sheets(cmp, mdl, results, qs)
    (cmp / "out").mkdir(exist_ok=True)
    try:
        TT.fill(tpl_path, SPEC, data, values(cmp, mdl, results, qs), book, kept={T_Q: ans.kept(BOOK).get(T_Q, {})})
    except TT.TemplateError as e:
        sys.exit(f"テンプレートが決まりどおりでない: {e}")
    mark_path(book).write_text(book_mark(results, qs) + "\n", encoding="utf-8")
    return book, said


def mark_path(book):
    """書き出したときの中身の印（点検が、いまのデータと比べる）"""
    return book.with_name(book.stem + ".mark")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cmd", nargs="?", choices=["default"])
    ap.add_argument("-w", "--work")
    ap.add_argument("--template")
    a = ap.parse_args()
    if a.cmd == "default":
        make_template(ASSETS / TEMPLATE)
        print(f"既定のテンプレートを作り直した → {ASSETS / TEMPLATE}")
        return
    if not a.work:
        ap.error("-w が要る")
    book, said = write(work_dir(a.work), a.template)
    for s in said:
        print(s)
    print(f"{book}")


if __name__ == "__main__":
    main()
