#!/usr/bin/env python3
"""selftest.py — 拠点間比較の点検が、仕込んだ誤りをちゃんと FAIL にするか・計算が手で数えた値と合うかを確かめる（LLM は使わない）。

  python3 evals/selftest.py                          # 架空の 3 拠点を一時フォルダに組み立てて試す（1〜2 分）
  python3 evals/selftest.py -w <フォルダ>             # 組み立てる場所を決める（残る。あれば組み立てを飛ばす）
  python3 evals/selftest.py -w <フォルダ> --only 用語  # 名前にその語を含む試験だけ

架空の 3 拠点（東京支店＝基準・詳しい・下位フローあり／大阪支店＝共通の資料の書き分けで支店長が承認・支払が粗い／
シンガポール拠点＝英語・申請と承認が 1 本・固有の作業 1 つ）を evals/fixture_tools.py で組み立て、
 (1) 正しいデータで全段 FAIL 0・WARN 0（拠点の門＝asis の点検も）・形の違反 0・schemas.md がスキーマどおり
 (2) 計算が、手で数えた期待値と同じ（区分・記号・差の種類・粒度・複雑性（原典・そろえた）・差の大きさ・追加の資料の依頼・番号）
 (3) 同じ入力で 2 回書き出して、計算・Excel・draw.io が同じ
 (4) 納品物の表記（拠点の列の見出し・英語の原文やフローの名前が出ない）
 (5) 基準なし（横に並べる）でも FAIL 0 で、差に関わる拠点の決め方が変わる
 (6) 誤りを 1 つずつ仕込んだ写しで、狙った FAIL が出る
 (7) 部品: 下書き・候補・番号の引き継ぎ・番号の振り直しの付け直し・回答の取り込み・進み具合・納品の門（draw.io の代わりの試験の道具で）
を確かめる。スクリプトを直したら、これを通してから使う。終わりの行が「NG 0」でなければ終了コード 1。
"""
import argparse
import json
import re
import os
import pathlib
import shutil
import subprocess
import sys
import tempfile

HERE = pathlib.Path(__file__).resolve().parent
SKILL = HERE.parent
S = SKILL / "scripts"
sys.path.insert(0, str(S))
sys.path.insert(0, str(HERE))

import fixture_tools as FT  # noqa: E402

NG = []
ONLY = None


def say(ok, name, detail=""):
    if not ok:
        NG.append(name)
    print(f"{'OK' if ok else 'NG'}  {name}" + ("" if ok else f" — {detail}"))


def wanted(name):
    return ONLY is None or ONLY in name


def run(script, *args, env=None):
    env = dict(env or os.environ, ASIS_NO_PREVIEW="1")               # draw.io の無い試験で、下見の描画（preview_mx）に時間を使わない
    r = subprocess.run([sys.executable, str(S / script), *map(str, args)], capture_output=True, text=True, env=env)
    return r.returncode, r.stdout + r.stderr


def check(cmp, stage=None, fast=True):
    return run("sc_check.py", "-w", cmp, *(["--stage", stage] if stage else []), *(["--fast"] if fast else []), "--all-lines")


def jload(p):
    return json.loads(pathlib.Path(p).read_text(encoding="utf-8"))


def jsave(p, d):
    pathlib.Path(p).write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")


def edit(p, fn):
    d = jload(p)
    fn(d)
    jsave(p, d)


def tail(out, n=6):
    return " / ".join([x for x in out.strip().splitlines() if x.startswith(("FAIL", "WARN"))][:n] or out.strip().splitlines()[-3:])


def fake_drawio(root, ok=True):
    """draw.io の代わり（試験だけ）: -o の次の道に、空の PNG を書く（ok=False は何も書かずに失敗する）"""
    b = root / ("bin_ok" if ok else "bin_ng")
    b.mkdir(exist_ok=True)
    p = b / "drawio"
    if ok:
        p.write_text("#!/bin/sh\nwhile [ $# -gt 0 ]; do if [ \"$1\" = \"-o\" ]; then "
                     f"\"{sys.executable}\" -c \"from PIL import Image; Image.new('RGB',(8,8),'white').save('$2')\"; fi; shift; done\n", encoding="utf-8")
    else:
        p.write_text("#!/bin/sh\necho 'export failed' >&2\nexit 1\n", encoding="utf-8")
    p.chmod(0o755)
    return dict(os.environ, PATH=f"{b}{os.pathsep}{os.environ.get('PATH', '')}")


# ---------- 組み立て ----------

def build_base(root):
    base = root / "base"
    if (base / "out" / "拠点比較フロー.drawio").exists():
        return base
    shutil.rmtree(base, ignore_errors=True)
    FT.build_sites(base, root / "base_in")
    FT.build_compare(base)
    for script in ("sc_book.py", "sc_drawio.py"):
        rc, out = run(script, "-w", base)
        if rc:
            sys.exit(f"{script} が失敗した:\n{out[-1500:]}")
    return base


def copy(base, root, name):
    d = root / "cases" / name
    shutil.rmtree(d, ignore_errors=True)
    shutil.copytree(base, d, symlinks=True)
    return d


# 道の名前
def P_flow(d, sid, fid):
    return d / "asis" / sid / "flows" / f"{fid}.flow.json"


def P_match(d, rid):
    return d / "rows" / f"{rid}.match.json"


def task(m, tid):
    return next(t for t in m["tasks"] if t["id"] == tid)


def node(fl, nid):
    return next(n for n in fl["nodes"] if n["id"] == nid)


def axis_site(ax, rid, sid):
    r = next(r for r in ax["rows"] if r["id"] == rid)
    return next(x for x in r["sites"] if x["site"] == sid)


def recompute(d):
    run("sc_compute.py", "-w", d)


def rebuild(d):
    recompute(d)
    run("sc_book.py", "-w", d)
    run("sc_drawio.py", "-w", d)


def side_by_side(d):
    """基準なしにする（sites.json の baseline を null に、シンガポール拠点を先頭に。比べる行は merged）"""
    def s(x):
        x["baseline"] = None
        x["sites"] = [x["sites"][2], x["sites"][1], x["sites"][0]]
    edit(d / "sites.json", s)

    def a(x):
        x["mode"] = "side_by_side"
        for r in x["rows"]:
            r["origin"] = "merged"
    edit(d / "axis.json", a)


# ---------- (1) 正しいデータ ----------

def good(base):
    rc, out = check(base, fast=False)
    say(rc == 0 and "0 FAIL / 0 WARN" in out, "正しいデータ: 全段（拠点の門を含む）FAIL 0・WARN 0", tail(out))
    rc, out = run("sc_validate.py", "-w", base)
    say(rc == 0 and "違反があるのは 0 ファイル" in out, "正しいデータ: 形の違反 0", out.strip()[-200:])
    rc, out = run("sc_validate.py", "--check-doc")
    say(rc == 0 and "合っている" in out, "schemas.md の項目表がスキーマと同じ", out.strip()[-200:])
    rc, out = check(base, "verify")
    say("支持率 100%" in out, "正しいデータ: 対応の照合の支持率 100%", tail(out))


# ---------- (2) 計算の期待値 ----------

def expected(base):
    res = {r: jload(base / "rows" / f"{r}.result.json") for r in ("B01", "B02", "B03")}
    T = {t["no"]: t for r in res.values() for t in r["tasks"]}
    kinds = lambda no: [(x["kind"], x["state"], x["sites"]) for x in T[no]["diffs"]]
    want = {
        "T-001": ("差異あり", {"TK": "○", "OS": "○", "SG": "—"}, [("記載なし", "記載なし", ["SG"])]),
        "T-002": ("差異あり", {"TK": "○", "OS": "○", "SG": "△"}, [("内容", "方法", ["SG"])]),
        "T-003": ("差異あり", {"TK": "○", "OS": "△", "SG": "○"}, [("担当", "違う", ["OS"])]),
        "T-004": ("拠点固有", {"TK": "○", "OS": "—", "SG": "—"}, [("記載なし", "記載なし", ["OS", "SG"])]),
        "T-005": ("拠点固有", {"TK": "—", "OS": "—", "SG": "○"}, [("記載なし", "記載なし", ["TK", "OS"])]),
        "T-006": ("共通", {"TK": "○", "OS": "○", "SG": "○"}, []),
        "T-007": ("差異あり", {"TK": "○", "OS": "—", "SG": "△"},
                  [("担当", "違う", ["SG"]), ("システム", "違う", ["SG"]), ("内容", "方法", ["SG"]), ("記載なし", "記載なし", ["OS"])]),
    }
    for no, (cls, marks, ds) in want.items():
        t = T.get(no)
        got = (t["class"], t["marks"], kinds(no)) if t else None
        say(got == (cls, marks, ds), f"計算: {no} の区分・記号・差の種類", f"{got}")
    say([t["no"] for t in res["B03"]["tasks"]] == ["T-004", "T-005", "T-006", "T-007"], "計算: 支払処理のタスクの並び（基準の箱の並び。固有の作業は前の基準のタスクのあと）",
        f"{[t['no'] for t in res['B03']['tasks']]}")
    g = T["T-006"]["grain"]
    say(g["counts"] == {"TK": 5, "OS": 1, "SG": 1} and g["coarse"] == ["OS", "SG"] and g["fine"] == ["TK"], "計算: 粒度（粗く書いた拠点と、内訳の数）", f"{g}")
    one = [(x["kind"], x["sites"]) for x in T["T-006"]["onesided"]]
    say(one == [("システム", ["OS", "SG"]), ("帳票", ["SG"])], "計算: 片方だけ記載は差に数えない（共通のまま）", f"{one}")
    c = res["B03"]["sites"]["TK"]["complexity"]
    say((c["raw"]["br"], c["raw"]["hand"], c["raw"]["score"], c["aligned"]["br"], c["aligned"]["hand"], c["aligned"]["score"]) == (1, 2, 4, 1, 0, 2),
        "計算: 東京支店の支払処理の複雑性（原典 分岐 1・受け渡し 2・点 4／そろえた 分岐 1・受け渡し 0・点 2）", f"{c}")
    c = res["B02"]["sites"]["OS"]["complexity"]
    say((c["raw"]["br"], c["raw"]["hand"], c["raw"]["score"], c["aligned"]["score"]) == (1, 1, 3, 0),
        "計算: 大阪支店の承認の複雑性（原典 課長→支店長の受け渡し 1。そろえた粒度では 1 つのタスクの中なので 0）", f"{c}")
    sp = {s: (v.get("spread") or {}).get("diff") for s, v in res["B03"]["sites"].items()}
    say(sp == {"TK": None, "OS": -2, "SG": -2} and res["B03"]["spread"] == "小", "計算: 差の大きさ（基準との点数の差）", f"{sp} {res['B03']['spread']}")
    rq = {q["site"]: (q["coarse"], q["absent"], q["gain"], q["folded"]["decisions"], q["folded"]["handoffs"]) for q in res["B03"]["requests"]}
    say(rq.get("OS") == (["T-006"], ["T-004", "T-005", "T-007"], 4, 1, 2) and rq.get("SG") == (["T-006"], ["T-004"], 2, 1, 2),
        "計算: 追加の資料の依頼（粗く書いたタスク・記載なし・活かせていない判断と受け渡し・比べられるようになる数）", f"{rq}")
    os_kinds = next(q["kinds"] for q in res["B03"]["requests"] if q["site"] == "OS")
    say(any("判断の基準" in k for k in os_kinds) and any("業務分掌" in k for k in os_kinds), "計算: 依頼する資料の種類（判断 → 規程・基準表、受け渡し → 業務分掌）", f"{os_kinds}")
    nb = jload(base / "numbers.json")
    say(nb["next"] == 8 and nb["assigned"]["B03|t4"] == "T-007" and nb["qassigned"].get("rows/B03|q1") == "Q-003", "番号: タスクの通し番号と確認事項の番号", f"{nb}")


def unit():
    import sc_common as C
    import sc_compute as K
    cfg = K.config()
    say([K.level_of(x, cfg) for x in (0, 4, 5, 10, 11)] == ["小", "小", "中", "中", "大"], "部品: 大中小の境（点 4 まで小・10 まで中）")
    say([K.spread_level(x, cfg) for x in (0, 2, 3, 5, 6)] == ["小", "小", "中", "中", "大"], "部品: 差の大きさの境（3 から中・6 から大）")
    say(K.compare_sets({"課長"}, {"課長"}, False, False) == "同じ" and K.compare_sets(set(), {"課長"}, False, False) == "片方だけ記載"
        and K.compare_sets({"課長"}, {"課長", "部長"}, True, False) == "同じ" and K.compare_sets({"課長"}, {"課長", "部長"}, False, False) == "違う",
        "部品: 担当・システム・帳票の比べ方（粗い側が含まれていれば同じ・書いていないのは片方だけ記載）")
    say(K.involved_sites([("TK", "OS")], ["TK", "OS", "SG"], "TK") == ["OS"] and K.involved_sites([("OS", "SG")], ["TK", "OS", "SG"], "TK") == ["OS", "SG"]
        and K.involved_sites([("TK", "OS"), ("SG", "OS")], ["TK", "OS", "SG"], None) == ["OS"],
        "部品: 差に関わる拠点（基準と違う拠点・基準で決まらない組は両方・基準なしは多くと違う拠点）")
    say(C.likeness("申請を承認する", "申請を確認して承認する") >= 0.34 and C.likeness("申請を承認する", "振込結果を自動で記帳する") < 0.34,
        "部品: 記載なしの候補の似ている度合い（かなを除いた字の組）")
    from sc_rows import linear_order
    nodes = ["F:start", "F:P-001", "F:D-001", "F:P-002", "F:P-003", "F:end"]
    edges = [("F:start", "F:P-001"), ("F:P-001", "F:D-001"), ("F:D-001", "F:P-002"), ("F:D-001", "F:P-003"), ("F:P-002", "F:P-001"), ("F:P-003", "F:end")]
    say(linear_order(nodes, edges) == ["F:start", "F:P-001", "F:D-001", "F:P-002", "F:P-003", "F:end"], "部品: 箱の並び（戻りの線は数えない・枝は線の順）",
        f"{linear_order(nodes, edges)}")
    import openpyxl
    import sc_book
    cells = {v for ws in openpyxl.load_workbook(C.ASSETS / sc_book.TEMPLATE).worksheets for row in ws.iter_rows(values_only=True) for v in row if isinstance(v, str)}
    miss = [k for k, text in sc_book.LEGEND if text not in cells]
    say(not miss, "部品: 既定のテンプレートの読み方が、いまの説明（sc_book.LEGEND）と同じ（違えば python3 scripts/sc_book.py default）", f"{miss}")


# ---------- (3)(4) 揺れ・表記 ----------

def stable(base, root):
    d = copy(base, root, "stable")
    before = {p.name: p.read_bytes() for p in (d / "rows").glob("*.result.json")}
    fig0, mark0 = (d / "out" / "拠点比較フロー.drawio").read_text(encoding="utf-8"), (d / "out" / "拠点比較.mark").read_text(encoding="utf-8")
    for script in ("sc_compute.py", "sc_book.py", "sc_drawio.py"):         # ハッシュの種を変えて（set の順に頼ると割れる）
        run(script, "-w", d, env=dict(os.environ, PYTHONHASHSEED="4"))
    after = {p.name: p.read_bytes() for p in (d / "rows").glob("*.result.json")}
    say(before == after, "揺れ: 計算し直しても（ハッシュの種を変えても）、計算の結果が同じ")
    say(fig0 == (d / "out" / "拠点比較フロー.drawio").read_text(encoding="utf-8"), "揺れ: 書き出し直しても、draw.io が同じ")
    say(mark0 == (d / "out" / "拠点比較.mark").read_text(encoding="utf-8"), "揺れ: 書き出し直しても、Excel の中身の印が同じ")
    shutil.rmtree(d)


def wording(base):
    import openpyxl
    wb = openpyxl.load_workbook(base / "out" / "拠点比較.xlsx")
    head = [c.value for c in wb["拠点タスク一覧"][1]]
    say(head[5:8] == ["【基準】東京支店", "大阪支店", "シンガポール拠点"], "表記: 拠点タスク一覧の拠点の列（基準が左・【基準】の印・拠点の名前）", f"{head}")
    cells = [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]
    fig = (base / "out" / "拠点比較フロー.drawio").read_text(encoding="utf-8")
    eng = ["Claim and approval", "submits", "reviews the claim", "｜Employee｜", "帳票: Receipt"]
    say(not [w for w in eng if any(w in v for v in cells) or w in fig], "表記: 英語の原文・英語のフローの名前・英語の担当と帳票の名前が、Excel と draw.io に出ない",
        f"{[w for w in eng if any(w in v for v in cells) or w in fig]}")
    say(any("申請と承認（申請の箱）" in v for v in cells), "表記: 日本語でない拠点のフローは、日本語の名前（title）で出る")
    say("｜申請者（Employee）｜経費精算書を領収書と一緒に提出する" in fig and "Expense Claim Procedure (Singapore) p.1" in fig
        and 'value="申請者（Employee）"' in fig and 'value="申請者"' in fig,
        "表記: 図のレーン・説明の担当は比べる用語の名前（日本語）に拠点の言葉を添える・出典の資料名は固有名のまま")
    say("T-006" in fig and "内訳" in fig and "記載なし" in fig, "表記: 図にタスクの番号・内訳の枠・記載なしの空き枠がある")


# ---------- (5) 基準なし ----------

def no_baseline(base, root):
    d = copy(base, root, "side_by_side")
    side_by_side(d)
    rc, out = run("sc_compute.py", "-w", d)
    rc2, out2 = check(d, "rows")
    rc3, out3 = check(d, "compute")
    say(rc == 0 and rc2 == 0 and rc3 == 0, "基準なし: 計算・タスクの対応・計算の点検が FAIL 0", out[-300:] + tail(out2) + tail(out3))
    res = {r: jload(d / "rows" / f"{r}.result.json") for r in ("B02", "B03")}
    t3 = res["B02"]["tasks"][0]
    say(t3["marks"] == {"SG": "○", "OS": "△", "TK": "△"},
        "基準なし: 承認の担当の差（課長／課長と支店長）は、粗く書いたシンガポール拠点がどちらとも同じなので、違う 2 拠点の両方に △", f"{t3['marks']}")
    t7 = next(t for t in res["B03"]["tasks"] if t["id"] == "t4")
    say(t7["marks"]["SG"] == "△" and t7["marks"]["TK"] == "△", "基準なし: 2 拠点だけの差（記帳の担当・システム）は、両方に △", f"{t7['marks']}")
    say(all(v.get("spread") is not None or v["complexity"]["aligned"]["score"] == min(x["complexity"]["aligned"]["score"] for x in res["B03"]["sites"].values() if x["relation"] != "none")
            for v in res["B03"]["sites"].values()), "基準なし: 差の大きさは、点数のいちばん小さい拠点と比べる")
    rc, out = run("sc_book.py", "-w", d)
    import openpyxl
    head = [c.value for c in openpyxl.load_workbook(d / "out" / "拠点比較.xlsx")["拠点タスク一覧"][1]]
    say(rc == 0 and head[5:8] == ["シンガポール拠点", "大阪支店", "東京支店"], "基準なし: 拠点の列は sites.json の順・【基準】の印なし", f"{head}")
    rc, out = run("sc_drawio.py", "-w", d)
    rc2, out2 = check(d, "outputs")
    say(rc == 0 and rc2 == 0, "基準なし: 書き出しと納品物の点検が FAIL 0", out[-300:] + tail(out2))
    shutil.rmtree(d)


# ---------- (6) 仕込んだ誤り ----------

def M(rid, fn):
    return lambda d: edit(P_match(d, rid), fn)


def F(sid, fid, fn):
    return lambda d: edit(P_flow(d, sid, fid), fn)


def J(name, fn):
    return lambda d: edit(d / name, fn)


def rm(*parts):
    return lambda d: (d.joinpath(*parts)).unlink()


def garbage(*parts):
    return lambda d: d.joinpath(*parts).write_text("{", encoding="utf-8")


def seq(*fns):
    def go(d):
        for f in fns:
            f(d)
    return go


def ax_row(rid, fn):
    return J("axis.json", lambda a: fn(next(r for r in a["rows"] if r["id"] == rid)))


def ax_site(rid, sid, fn):
    return J("axis.json", lambda a: fn(axis_site(a, rid, sid)))


def term(gid, fn):
    return J("terms.json", lambda t: fn(next(g for g in t["groups"] if g["id"] == gid)))


def drop_term(gid):
    return J("terms.json", lambda t: t.update(groups=[g for g in t["groups"] if g["id"] != gid]))


def applies(key, fn):
    return J("applies.json", lambda a: fn(next(i for i in a["items"] if i["key"] == key)))


def shared_card(d):
    edit(d / "asis" / "TK" / "cards" / "M01" / "p0002.json", lambda c: c.update(summary=c.get("summary", "") + "（東京で直した）"))


def stale_snap(d):
    edit(P_flow(d, "TK", "F01"), lambda f: node(f, "P-001").update(label="経費申請書を作成して保存する"))


def no_sources(d):
    edit(P_flow(d, "TK", "F01"), lambda f: node(f, "P-001").pop("sources"))


def verify_report(rid, fn):
    def go(d):
        edit(d / "reports" / f"verify_{rid}.json", fn)
        run("sc_verify.py", "draft", "-w", d)
    return go


def finding(d):
    verify_report("B03", lambda r: r["findings"].append({"kind": "対応違い", "target": "t1", "text": "一覧の出力は、支払データの作成の一部に読める"}))(d)


def review_finding(**kw):
    return seq(finding, J("review/B03.review.json", lambda r: r["findings"][0].update(**kw)))


def all_unsupported(r):
    for v in r["verdicts"]:
        v.update(verdict="不支持", evidence="同じ作業に読めない")


def book_cell(d):
    import openpyxl
    p = d / "out" / "拠点比較.xlsx"
    wb = openpyxl.load_workbook(p)
    wb["概要"]["A60"] = "M01 の 2 ページを見る"
    wb.save(p)


def fig_hand(d):
    p = d / "out" / "拠点比較フロー.drawio"
    p.write_text(p.read_text(encoding="utf-8").replace("経費申請書を作成する", "経費申請書を作る", 1), encoding="utf-8")


def overlap(d):
    def go(lay):
        pg = next(p for p in lay["pages"] if p["id"] == "row-B01")
        pg["nodes"]["TK.F01.P-002"] = list(pg["nodes"]["TK.F01.P-001"])
    edit(d / "out" / "拠点比較フロー.layout.json", go)


def new_question(**kw):
    q = {"id": "q9", "category": "用語", "effect": "差の種類", "question": "経理課長は経理部の中の担当ですか、別の担当ですか。", "assumption": "別の組にした", "ask": "東京支店の経理部"}
    q.update(kw)
    return q


def extra_site(d):
    edit(d / "sites.json", lambda s: s["sites"].append({"id": "HK", "name": "香港拠点", "aliases": ["香港"], "language": "en", "inputs": [str(d / "nowhere")]}))


def snap_x(sid, key, fn):
    return J(f"snap/{sid}.json", lambda s: fn(next(x for x in s["excluded"] if x["key"] == key)))


def os_excludes_own(d):
    """大阪支店が、自分に当たる書き分け（大阪支店だけの承認 M01.2.s4）を、ほかの拠点だけの手順として外したことにする"""
    def go(s):
        x = next(x for x in s["excluded"] if x["key"] == "M01.2.s3")
        s["excluded"].append(dict(x, key="M01.2.s4", text="申請を承認する（支店長）"))
    edit(d / "snap" / "OS.json", go)


def split_b02(m):
    """B02: 東京支店の課長の承認（大阪支店では行わない手順）を、別のタスクに分ける"""
    task(m, "t1")["members"]["TK"].remove("F02:P-003")
    m["tasks"].append({"id": "t2", "name": "課長が申請を承認する", "members": {"TK": ["F02:P-003"], "OS": [], "SG": []},
                       "absent": {"OS": {"checked": [], "note": "大阪支店は支店長が承認する"}, "SG": {"checked": [], "note": "課長の承認だけを分けて書いていない"}},
                       "joined": "", "content": []})


def other_finance(t):
    """シンガポール拠点の Finance を、東京支店の担当（site）の組に分ける（シンガポール拠点の経理の箱が、ほかの拠点の作業になる）"""
    g = next(g for g in t["groups"] if g["id"] == "R03")
    g["members"] = [m for m in g["members"] if m["site"] != "SG"]
    t["groups"].append({"id": "R09", "kind": "actor", "name": "東京支店の経理部", "members": [{"site": "SG", "term": "Finance"}],
                        "basis": "same_name", "site": "TK", "reason": "シンガポール拠点の手順書の Finance は、東京支店の経理部"})


def to_sg(s):
    """東京支店・大阪支店の「経費申請書を作成する」を、シンガポール拠点へ行う作業にする（シンガポール拠点は受け取る側）"""
    s["nodes"]["F01:P-001"]["label"] = "経費申請書を作成してシンガポール拠点へ送付する"


def BD_units(s):
    import build_drawio
    return build_drawio.units(s)


def split_task(m):
    t = task(m, "t3")
    t["members"]["TK"].remove("F03-1:P-002")
    m["tasks"].append({"id": "t5", "name": "口座を確認して再振込する", "members": {"TK": ["F03-1:P-002"], "OS": [], "SG": []},
                       "absent": {"OS": {"checked": [], "note": "再振込は書いていない"}, "SG": {"checked": [], "note": "再振込は書いていない"}},
                       "joined": "", "content": []})


CASES = [
    # [sites]
    ("拠点 ID の重なり", "sites", J("sites.json", lambda s: s["sites"][2].update(id="OS")), "拠点 ID が重なっている"),
    ("基準の拠点が一覧に無い", "sites", J("sites.json", lambda s: s.update(baseline="XX")), "が、拠点の一覧に無い"),
    ("別名が 2 つの拠点に", "sites", J("sites.json", lambda s: s["sites"][1]["aliases"].append("東京")), "の両方の名前・別名になっている"),
    ("sites.json の余分な項目", "sites", J("sites.json", lambda s: s.update(memo="x")), "sites.json: 形が違う"),
    ("拠点の asis の作業フォルダが無い", "sites", extra_site, "asis の作業フォルダができていない"),
    ("拠点の門: 出典の無い箱（asis の点検が FAIL）", "sites", no_sources, "拠点の門を通らない", False),
    ("共通の資料のカードを拠点の側で直した", "sites", shared_card, "のカードが、共通の作業フォルダのものと違う"),
    ("共通の資料の読取が終わっていない", "sites", rm("asis", "shared", "cards", "M01", "p0002.json"), "の読取が終わっていない"),
    ("共通の資料をページ化していない", "sites", J("asis/shared/manifest.json", lambda m: m["docs"][0].update(rendered=False)), "をページ化していない"),
    ("共通の資料の棚卸しが無い", "sites", rm("asis", "shared", "manifest.json"), "共通の資料の棚卸しが無い"),
    ("共通の資料を使う拠点の棚卸しに、共通の資料が無い", "sites", J("sites.json", lambda s: s["sites"][2].pop("uses_shared")), "が、拠点の棚卸しに入っていない"),
    ("共通の資料が変わった（拠点の棚卸しと中身が違う）", "sites", J("asis/TK/manifest.json", lambda m: m["docs"][0].update(sha1="0" * 40)), "の中身が、共通の作業フォルダの棚卸しと違う"),
    # [applies]
    ("書き分けが無い", "applies", rm("applies.json"), "applies.json が無い"),
    ("書き分けが読めない", "applies", garbage("applies.json"), "applies.json が読めない"),
    ("書き分けの余分な項目", "applies", applies("M01.2.s1", lambda i: i.update(memo="x")), "applies.json: 形が違う"),
    ("拾った項目が書き分けに無い", "applies", J("applies.json", lambda a: a.update(items=[i for i in a["items"] if i["key"] != "M01.2.s4"])), "拾った項目 M01.2.s4"),
    ("書き分けの資料が変わった", "applies", applies("M01.2.s1", lambda i: i.update(doc="0" * 40)), "共通の資料の中身が変わった"),
    ("書き分けが未定", "applies", applies("M01.2.s1", lambda i: i.update(scope="未定")), "当てはまる拠点が未定"),
    ("only なのに拠点が無い", "applies", applies("M01.2.s4", lambda i: i.update(sites=[])), "only なのに拠点（sites）が無い"),
    ("all なのに拠点がある", "applies", applies("M01.2.s1", lambda i: i.update(sites=["OS"])), "all なのに拠点（sites）が書いてある"),
    ("拠点を絞った根拠が無い", "applies", applies("M01.2.s4", lambda i: i.update(evidence=" ")), "拠点を絞った根拠（evidence）が無い"),
    ("書き分けの拠点が sites.json に無い", "applies", applies("M01.2.s4", lambda i: i.update(sites=["XX"])), "拠点 XX が sites.json に無い"),
    ("大阪支店だけの手順が東京支店のフローに入る", "applies",
     F("TK", "F02", lambda f: node(f, "P-003")["sources"].append({"doc_id": "M01", "page": 2, "quote": "※ 大阪支店では、課長に代えて支店長が申請を承認する。",
                                                                    "from_image": False, "item": "s4"})), "ほかの拠点だけの手順「申請を承認する"),
    ("ほかの拠点だけの手順を外す理由が決まった形でない", "applies", F("TK", "F02", lambda f: f["excluded_items"][0].update(reason="対象外")), "外す理由が、決まった形でない"),
    ("大阪支店が自分の手順を外す", "applies",
     F("OS", "F02", lambda f: f["excluded_items"].append({"doc_id": "M01", "page": 2, "item": "s4", "reason": "ほかの拠点だけの手順（東京支店だけ）"})),
     "を、ほかの拠点だけの手順として外している"),
    ("拠点で分かれる判断を残す", "applies", F("TK", "F02", lambda f: node(f, "D-001").update(label="大阪支店か")), "拠点で分かれる判断を置いている"),
    # [snapshot]
    ("拠点の写しが無い", "snapshot", rm("snap", "OS.json"), "拠点の写しが無い"),
    ("拠点の写しの形が崩れている", "snapshot", J("snap/OS.json", lambda s: s.update(site="TK")), "拠点の写しの形が崩れている"),
    ("拠点を直したあとの古い写し", "snapshot", stale_snap, "写したあとに変わった"),
    # [terms]
    ("用語の対応が無い", "terms", rm("terms.json"), "terms.json が無い"),
    ("用語の対応が読めない", "terms", garbage("terms.json"), "terms.json が読めない"),
    ("用語の対応の余分な項目", "terms", term("R01", lambda g: g.update(memo="x")), "terms.json: 形が違う"),
    ("用語の対応の形の違反が多い（畳む）", "terms", J("terms.json", lambda t: [g.update(memo="x") for g in t["groups"]]), "形の違反が、ほかに"),
    ("用語の組の ID の重なり", "terms", term("R02", lambda g: g.update(id="R01")), "組の ID R01 が重なっている"),
    ("用語の組の ID の頭が種類と合わない", "terms", term("R04", lambda g: g.update(id="Y09")), "ID の頭が種類と合わない"),
    ("担当の組の名前が日本語でない", "terms", term("R03", lambda g: g.update(name="Finance")), "担当・帳票の組の名前は日本語にする"),
    ("用語の組の拠点が sites.json に無い", "terms", term("R03", lambda g: g["members"].append({"site": "XX", "term": "経理"})), "拠点 XX が sites.json に無い"),
    ("1 つの用語が 2 つの組に", "terms", term("Y02", lambda g: g["members"].append({"site": "TK", "term": "経費精算システム"})), "の 2 つの組に入っている"),
    ("同じ拠点の 2 つの用語を 1 つの組に", "terms", seq(term("R02", lambda g: g["members"].append({"site": "TK", "term": "経理課長"})), drop_term("R04")),
     "を 1 つの組にしている"),
    ("名前が違う用語を same_name で", "terms", seq(term("Y01", lambda g: g["members"].append({"site": "SG", "term": "Oracle"})), drop_term("Y04")),
     "名前が違う用語を same_name でまとめている"),
    ("訳でまとめた理由が無い", "terms", term("R01", lambda g: g.update(reason="")), "でまとめた理由（reason）が無い"),
    ("比べる箱の用語がどの組にも無い", "terms", drop_term("R05"), "「支店長」がどの組にも無い"),
    ("用語の確認事項の組が無い", "terms", J("terms.json", lambda t: t["questions"].append(new_question(group="R99"))), "組 R99 が無い"),
    ("ほかの拠点の担当の拠点が sites.json に無い", "terms", term("R05", lambda g: g.update(site="XX", reason="x")), "site の拠点 XX が sites.json に無い"),
    ("ほかの拠点の担当の名前なのに site が無い", "terms",
     J("terms.json", lambda t: t["groups"].append({"id": "R09", "kind": "actor", "name": "東京支店の経理部", "members": [{"site": "OS", "term": "東京支店の経理部"}],
                                                    "basis": "same_name", "reason": ""})),
     "は、ほかの拠点（東京支店）の担当の名前に読める"),
    ("site を担当でない組に書く", "terms", term("Y02", lambda g: g.update(site="TK")), "site は担当（actor）の組だけに書く"),
    ("site を書いた組に根拠が無い", "terms", term("R05", lambda g: g.update(site="TK", reason="")), "その担当が東京支店の人・部門である根拠を書く"),
    ("site の組に、その拠点自身の用語", "terms", term("R02", lambda g: g.update(site="TK", reason="x")), "東京支店自身の用語「課長」を入れない"),
    # [axis]
    ("比べる行が無い", "axis", rm("axis.json"), "axis.json が無い"),
    ("比べる行が読めない", "axis", garbage("axis.json"), "axis.json が読めない"),
    ("比べる行の余分な項目", "axis", ax_row("B01", lambda r: r.update(memo="x")), "axis.json: 形が違う"),
    ("基準ありなのに side_by_side", "axis", J("axis.json", lambda a: a.update(mode="side_by_side")), "mode が side_by_side"),
    ("業務分類の ID の重なり", "axis", J("axis.json", lambda a: a["groups"].append({"id": "G01", "name": "経費", "name_from": None})), "業務分類の ID が重なっている"),
    ("業務分類の名前が日本語でない", "axis", J("axis.json", lambda a: a["groups"][0].update(name="Expenses")), "業務分類の名前は日本語にする"),
    ("業務分類の名前の出どころが写しに無い", "axis", J("axis.json", lambda a: a["groups"][0].update(name_from={"site": "TK", "category": "F99"})), "name_from の分類 F99 が"),
    ("行の ID の重なり", "axis", ax_row("B02", lambda r: r.update(id="B01")), "行の ID が重なっている"),
    ("行の業務分類が無い", "axis", ax_row("B01", lambda r: r.update(group="G09")), "業務分類 G09 が groups に無い"),
    ("親の行が無い", "axis", ax_row("B01", lambda r: r.update(parent="B09")), "親の行 B09 が無い"),
    ("行の名前が日本語でない", "axis", ax_row("B01", lambda r: r.update(name="Expense claim")), "行の名前は日本語にする"),
    ("行の名前の出どころのフローが無い", "axis", ax_row("B01", lambda r: r["name_from"].update(flow="F09")), "name_from のフロー F09 が"),
    ("行に拠点の当て方が無い", "axis", ax_row("B01", lambda r: r.update(sites=[x for x in r["sites"] if x["site"] != "OS"])), "大阪支店の当て方が 0 回ある"),
    ("origin が baseline なのに基準のフローが無い", "axis", ax_site("B01", "TK", lambda x: x.update(relation="none", parts=[], reason="無い")),
     "origin が baseline なのに"),
    ("not_in_baseline なのに基準のフローを当てる", "axis", ax_row("B01", lambda r: r.update(origin="not_in_baseline")), "origin が not_in_baseline なのに"),
    ("基準ありで merged", "axis", ax_row("B01", lambda r: r.update(origin="merged")), "基準ありでは、origin は baseline か not_in_baseline"),
    ("基準なしで merged でない", "axis", seq(side_by_side, ax_row("B01", lambda r: r.update(origin="baseline"))), "基準なしでは、origin は merged"),
    ("行の拠点が sites.json に無い", "axis", ax_row("B01", lambda r: r["sites"].append({"site": "XX", "relation": "none", "parts": [], "reason": "無い"})),
     "拠点 XX が sites.json に無い"),
    ("same なのにフローが無い", "axis", ax_site("B01", "OS", lambda x: x.update(parts=[])), "same なのに、当てるフロー（parts）が無い"),
    ("none なのにフローを当てる", "axis", ax_site("B01", "OS", lambda x: x.update(relation="none", reason="無い")), "none なのに、フローを当てている"),
    ("split の理由が無い", "axis", ax_site("B01", "SG", lambda x: x.update(reason="")), "split の理由（reason）が無い"),
    ("当てたフローが写しに無い", "axis", ax_site("B03", "OS", lambda x: x["parts"][0].update(flow="F09")), "フロー F09 が写しに無い"),
    ("日本語でない拠点のフローに日本語の名前が無い", "axis", ax_site("B03", "SG", lambda x: x["parts"][0].pop("title")), "日本語の名前（title）を書く"),
    ("行に入れた箱がフローに無い", "axis", ax_site("B01", "SG", lambda x: x["parts"][0].update(nodes=["F01:P-009"])), "箱 F01:P-009 が、フロー F01"),
    ("1 つの箱が 2 つの行に", "axis", ax_site("B02", "SG", lambda x: x["parts"][0]["nodes"].append("F01:P-001")), "の 2 か所に入っている"),
    ("基準に無い作業の箱を落とす（どの行にも無い）", "axis", ax_site("B02", "SG", lambda x: x["parts"][0].update(nodes=["F01:SUB-001"])), "がどの行にも入っていない"),
    ("行の材料が上限を超える", "axis", J("snap/TK.json", lambda s: s["nodes"]["F03:P-002"].update(detail="支払データの項目を確かめる。" * 3200)),
     "タスクの対応の材料が"),
    ("比べないフローが写しに無い", "axis", J("axis.json", lambda a: a["not_compared"].append({"site": "TK", "flow": "F09", "nodes": None, "reason": "総則"})),
     "not_compared: 東京支店のフロー F09 が写しに無い"),
    # [rows]
    ("タスクの対応が無い", "rows", rm("rows", "B01.match.json"), "rows/B01.match.json が無い"),
    ("タスクの対応が読めない", "rows", garbage("rows", "B01.match.json"), "rows/B01.match.json が読めない"),
    ("タスクの対応の余分な項目", "rows", M("B01", lambda m: task(m, "t1").update(memo="x")), "rows/B01.match.json: 形が違う"),
    ("row がファイルの名前と違う", "rows", M("B01", lambda m: m.update(row="B02")), "row が B02"),
    ("タスクの ID の重なり", "rows", M("B03", lambda m: task(m, "t2").update(id="t1")), "タスクの ID が重なっている"),
    ("タスクの名前が動詞で終わらない", "rows", M("B01", lambda m: task(m, "t1").update(name="経費申請書の作成")), "日本語の「目的語＋動詞」で終える"),
    ("行に当てていない拠点の箱がある", "rows", ax_site("B01", "SG", lambda x: x.update(relation="none", parts=[], reason="無い")), "は、この行にフローを当てていない（axis.json）のに、箱がある"),
    ("拠点の欄が無い", "rows", M("B01", lambda m: (task(m, "t1")["members"].pop("SG"), task(m, "t1")["absent"].pop("SG"))), "シンガポール拠点の欄が無い"),
    ("行に当てていない箱をタスクに入れる", "rows", M("B01", lambda m: task(m, "t1")["members"]["TK"].append("F02:P-001")), "は、この行に当てた箱でない"),
    ("作業でない箱をタスクに入れる", "rows", M("B02", lambda m: task(m, "t1")["members"]["TK"].append("F02:SUB-001")), "は作業の箱でない"),
    ("どの拠点の箱も無いタスク", "rows",
     M("B01", lambda m: m["tasks"].append({"id": "t3", "name": "申請を保存する", "members": {"TK": [], "OS": [], "SG": []},
                                         "absent": {s: {"checked": [], "note": "無い"} for s in ("TK", "OS", "SG")}, "joined": "", "content": []})),
     "どの拠点の箱も入っていない"),
    ("どの拠点も詳しいのに joined が空", "rows",
     M("B02", lambda m: (task(m, "t1")["members"].update(SG=[]), task(m, "t1")["absent"].update(SG={"checked": [], "note": "無い"}))),
     "どの拠点も 2 つ以上の箱でこのタスクを書いている"),
    ("箱があるのに記載なし", "rows", M("B01", lambda m: task(m, "t2")["absent"].update(SG={"checked": [], "note": "無い"})), "は箱があるのに、absent（記載なし）がある"),
    ("記載なしの根拠が無い", "rows", M("B03", lambda m: task(m, "t4")["absent"].pop("OS")), "に箱が無いのに、記載なしの根拠（absent）が無い"),
    ("記載なしで機械の候補を確かめていない", "rows", M("B03", lambda m: task(m, "t1")["absent"]["OS"].update(checked=[])), "を確かめていない"),
    ("記載なしの拠点が行に無い", "rows", M("B01", lambda m: task(m, "t1")["absent"].update(XX={"checked": [], "note": "無い"})), "は、この行にフローを当てていない"),
    ("記載なしの拠点を内容の差に書く", "rows", M("B03", lambda m: task(m, "t4")["content"][0].update(sites=["OS"])), "このタスクに箱が無い（記載なしは、内容の差に書かない）"),
    ("not_here の鍵が外した項目にも箱にも無い", "rows", M("B03", lambda m: task(m, "t1")["absent"]["OS"].update(not_here=["M01.9.s9"])),
     "は、この拠点がフローから外した項目にも、この拠点の箱にも無い"),
    ("not_here の項目を外した理由が書き分けでない", "rows",
     seq(snap_x("OS", "M01.2.s3", lambda x: x.update(reason="経費申請と同じなので描かない")), M("B03", lambda m: task(m, "t1")["absent"]["OS"].update(not_here=["M01.2.s3"]))),
     "の外した理由が「ほかの拠点だけの手順」でない"),
    ("not_here の項目が、書き分けでこの拠点に当たる", "rows",
     seq(os_excludes_own, M("B03", lambda m: task(m, "t1")["absent"]["OS"].update(not_here=["M01.2.s4"], checked=["F02:P-003", "M01.2.s3", "M01.2.s4"]))),
     "は、書き分け（applies.json）で大阪支店に当てはまる手順"),
    ("not_here の項目を checked に書かない", "rows", M("B03", lambda m: task(m, "t1")["absent"]["OS"].update(not_here=["M01.2.s3"], checked=["F02:P-003"])),
     "の not_here の M01.2.s3 を、checked にも書く"),
    ("not_here の箱に、ほかの拠点の名前が無い", "rows", M("B03", lambda m: task(m, "t1")["absent"]["OS"].update(not_here=["F01:P-001"])),
     "の名前・補足に、ほかの拠点の名前が無い"),
    ("書き分けで分かる記載なしに not_here が無い", "rows", M("B02", split_b02), "absent.not_here と checked に M01.2.s3 を書く"),
    ("not_here の受け取る側の箱が、この拠点へ行うと書いていない", "rows", M("B01", lambda m: task(m, "t1")["absent"]["SG"].update(not_here=["TK/F01:P-001"])),
     "の名前・補足に、シンガポール拠点へ行う作業とない"),
    ("not_here の受け取る側の箱が、このタスクの箱でない", "rows", M("B01", lambda m: task(m, "t1")["absent"]["SG"].update(not_here=["TK/F02:P-003"])),
     "は、このタスクの東京支店の箱でない"),
    ("not_here の受け取る側の拠点が、この拠点", "rows", M("B01", lambda m: task(m, "t1")["absent"]["SG"].update(not_here=["SG/F01:P-001"])),
     "の拠点 SG が、ほかの拠点でない"),
    ("受け取る側が機械で分かるのに not_here が無い", "rows", seq(J("snap/TK.json", to_sg), J("snap/OS.json", to_sg)),
     "どれもシンガポール拠点へ行う作業と書いている（シンガポール拠点は受け取る側）。absent.not_here に TK/F01:P-001, OS/F01:P-001 を書く"),
    ("対象外の拠点を内容の差に書く", "rows", J("terms.json", other_finance), "は、このタスクでは対象外"),
    ("内容の差の根拠に箱と項目の両方", "rows", M("B01", lambda m: task(m, "t2")["content"][0]["evidence"][0].update(item="M01.2.s3")),
     "根拠は 1 つずつ、node（箱）か item"),
    ("内容の差の根拠の項目が外した項目に無い", "rows", M("B01", lambda m: task(m, "t2")["content"][0]["evidence"].append({"site": "OS", "item": "M01.9.s9"})),
     "根拠の項目 M01.9.s9 が、大阪支店のフローから外した項目に無い"),
    ("内容の差の根拠の箱がタスクに無い", "rows", M("B03", lambda m: task(m, "t4")["content"][0]["evidence"][0].update(node="F03:P-001")), "根拠の箱 F03:P-001 が"),
    ("内容の差が日本語でない", "rows", M("B03", lambda m: task(m, "t4")["content"][0].update(text="Manual entry in Oracle")), "日本語で書く"),
    ("作業の箱がどのタスクにも無い", "rows", M("B03", lambda m: task(m, "t3")["members"]["TK"].remove("F03-1:P-002")), "がどのタスクにも入っていない"),
    ("1 つの箱が 2 つのタスクに", "rows", M("B03", lambda m: task(m, "t1")["members"]["TK"].append("F03:P-002")), "の 2 つのタスクに入っている"),
    ("確認事項のタスクが無い", "rows", M("B03", lambda m: m["questions"][0].update(task="t9")), "タスク t9 が無い"),
    ("確認事項の拠点が無い", "rows", M("B03", lambda m: m["questions"][0].update(site="XX")), "拠点 XX が無い"),
    # [compute]
    ("計算の結果が無い", "compute", rm("rows", "B01.result.json"), "rows/B01.result.json が無い"),
    ("計算の結果を手で直す", "compute", J("rows/B01.result.json", lambda r: r["tasks"][0].update({"class": "共通"})), "が、いまのデータと違う"),
    ("タスクを直して計算し直していない", "compute", M("B01", lambda m: task(m, "t1").update(name="経費申請書を作って保存する")), "が、いまのデータと違う"),
    # [verify]
    ("対応の照合をしていない", "verify", rm("review", "B01.review.json"), "対応の照合をしていない"),
    ("採否表の余分な項目", "verify", J("review/B01.review.json", lambda r: r.update(memo="x")), "review/B01.review.json: 形が違う"),
    ("採否表の判定を手で直す", "verify", J("review/B01.review.json", lambda r: r["verdicts"][0].update(verdict="一部支持")), "判定・指摘が、照合の報告と違う"),
    ("照合のあとにタスクを変える", "verify", seq(M("B01", lambda m: task(m, "t2").update(name="領収書を添えて申請書を出す")), recompute),
     "照合していない・照合したあとに変わった主張"),
    ("指摘に採否が無い", "verify", finding, "に採否（decision）と対応（action）が無い"),
    ("採用した指摘を反映していない", "verify", review_finding(decision="採用", action="直した"), "を、まだ反映していない"),
    ("採否の理由が無い", "verify", review_finding(decision="不採用", action=""), "対応・理由（action）が無い"),
    ("支持率 95% 未満", "verify", verify_report("B03", all_unsupported), "対応の照合の支持率が"),
    # [outputs]
    ("Excel が無い", "outputs", rm("out", "拠点比較.xlsx"), "拠点比較.xlsx が無い"),
    ("計算し直して Excel を書き出していない", "outputs", seq(M("B01", lambda m: task(m, "t1").update(name="経費申請書を作って保存する")), recompute),
     "拠点比較.xlsx が、いまのデータと違う"),
    ("Excel に資料 ID", "outputs", book_cell, "資料 IDが出ている"),
    ("draw.io が無い", "outputs", rm("out", "拠点比較フロー.drawio"), "拠点比較フロー.drawio が無い"),
    ("draw.io を手で直す", "outputs", fig_hand, "拠点比較フロー.drawio が、いまのデータと違う"),
    ("タスクの名前に箱の鍵（図）", "outputs", seq(M("B03", lambda m: task(m, "t1").update(name="F03:P-001 の一覧を出力する")), rebuild), "図の文字に箱の鍵が出ている"),
    ("タスクの名前に箱の鍵（Excel）", "outputs", seq(M("B03", lambda m: task(m, "t1").update(name="F03:P-001 の一覧を出力する")), rebuild), "箱の鍵が出ている（「"),
    ("英語の文が納品物に出る（行の当て方の理由）", "outputs", seq(ax_site("B01", "SG", lambda x: x.update(reason="申請と承認を 1 本（Claim and approval）に書いている")), rebuild),
     "英語が出ている（「Claim and approval」"),
    ("英語の文が図に出る（内容の差）", "outputs",
     seq(M("B03", lambda m: task(m, "t4")["content"][0].update(text="シンガポール拠点は records the payment in Oracle と書く")), rebuild), "図の文字に英語が出ている"),
    ("内容の差に行の ID", "outputs", seq(M("B03", lambda m: task(m, "t4")["content"][0].update(text="B03 の記帳は、東京支店は自動で、シンガポール拠点は経理が記録する")), rebuild),
     "行の IDが出ている"),
    ("図の箱が重なる", "outputs", overlap, "拠点比較フロー.drawio［経費申請］"),
]


def cases(base, root):
    for c in CASES:
        name, stage, fn, want = c[:4]
        fast = c[4] if len(c) > 4 else True
        if not wanted(name):
            continue
        d = copy(base, root, "case")
        try:
            fn(d)
        except Exception as e:                                       # 仕込み方の誤り（試験の誤り）
            say(False, f"仕込んだ誤り: {name}", f"仕込めない（{type(e).__name__}: {e}）")
            continue
        rc, out = check(d, stage, fast)
        say(rc == 1 and want in out and "形が崩れていて" not in out, f"仕込んだ誤り: {name}", f"「{want}」が出ない（終了 {rc}）: {tail(out)}")
    shutil.rmtree(root / "cases" / "case", ignore_errors=True)


# ---------- (7) 部品・進み具合・門 ----------

def parts(base, root):
    d = copy(base, root, "parts")
    if wanted("下書き"):
        rc, out = run("sc_init.py", root / "base_in", "-w", root / "parts_init", "--baseline", "東京支店")
        s = jload(root / "parts_init" / "sites.json") if (root / "parts_init" / "sites.json").exists() else {}
        say(rc == 0 and len(s.get("sites", [])) == 3 and s.get("shared") and s["baseline"] == next((x["id"] for x in s["sites"] if x["name"] == "東京支店"), None),
            "下書き: 拠点の一覧（入力のフォルダの形から 3 拠点・共通の資料・基準）", out[-300:])
        shutil.rmtree(root / "parts_init", ignore_errors=True)
        (d / "applies.json").unlink()
        rc, out = run("sc_applies.py", "draft", "-w", d)
        ap = jload(d / "applies.json")
        say(rc == 0 and {i["key"] for i in ap["items"]} == {"M01.2.s1", "M01.2.s2", "M01.2.s3", "M01.2.s4", "M01.2.d1"} and all(i["scope"] == "未定" for i in ap["items"])
            and next(i for i in ap["items"] if i["key"] == "M01.2.s4")["mentions"] == ["OS"],
            "下書き: 書き分けの候補（拠点の名前の出る項目と、同じ節の項目。未定）", f"{[(i['key'], i['scope'], i['mentions']) for i in ap['items']]}")
        shutil.copy2(base / "applies.json", d / "applies.json")
        rc, out = run("sc_site_ops.py", "-w", d, "--site", "OS")
        say(rc == 0 and '"op": "exclude", "item": "M01.2.s3"' in out and "ほかの拠点だけの手順（大阪支店では行わない" in out,   # 理由の頭（書き分けの根拠の文が続く）
            "下書き: ほかの拠点だけの手順を外す直し（大阪支店は、課長の承認を外す）", out[-300:])
        (d / "axis.json").rename(d / "axis.keep")
        rc, out = run("sc_axis.py", "draft", "-w", d)
        ax = jload(d / "axis.json")
        say(rc == 0 and [r["name"] for r in ax["rows"]] == ["経費申請", "承認", "支払処理"]
            and all(axis_site(ax, r["id"], "OS")["relation"] == "none" for r in ax["rows"]),
            "下書き: 比べる行（基準の分類から 3 行。ほかの拠点は none）", f"{[r['name'] for r in ax['rows']]}")
        (d / "axis.keep").replace(d / "axis.json")
        P_match(d, "B03").rename(d / "B03.keep")
        rc, out = run("sc_match.py", "draft", "-w", d, "--row", "B03")
        m = jload(P_match(d, "B03"))
        say(rc == 0 and len(m["tasks"]) == 7 and all(t["members"]["OS"] == [] for t in m["tasks"]) and m["tasks"][4]["name"] == "振込エラーがあるかを判断する",
            "下書き: タスク（基準の作業の箱 1 つ＝1 タスク。開いた下位フローの箱も）", f"{[(t['name'], t['members']['TK']) for t in m['tasks']]}")
        (d / "B03.keep").replace(P_match(d, "B03"))
        (d / "terms.json").rename(d / "terms.keep")
        rc, out = run("sc_terms.py", "draft", "-w", d)
        tm = jload(d / "terms.json")
        fin = [g for g in tm["groups"] if {"site": "SG", "term": "Finance"} in g["members"]]
        say(rc == 0 and len(fin) == 1 and fin[0]["basis"] == "same_name" and len(fin[0]["members"]) == 1,
            "下書き: 用語の対応（名前が同じものだけをまとめる。訳は LLM が）", f"{fin}")
        (d / "terms.keep").replace(d / "terms.json")
    if wanted("候補"):
        rc, out = run("sc_match.py", "candidates", "-w", d, "--row", "B03")
        say(rc == 0 and "F02:P-003" in out and "M01.2.s3" in out, "候補: 記載なしの候補（ほかの行の似た作業・フローから外した手順）", out[-400:])
        snap = d / "snap" / "OS.json"
        keep = snap.read_text(encoding="utf-8")
        sp = jload(snap)
        x = next(x for x in sp["excluded"] if x["key"] == "M01.2.s3")
        sp["excluded"] += [dict(x, flow="F99"), dict(x, key="M01.2.s9", reason="「支払処理」（F03）のフローで描く")]
        jsave(snap, sp)
        rc, out = run("sc_match.py", "candidates", "-w", d, "--row", "B03")
        lines = [ln for ln in out.splitlines() if "checked:" in ln]
        say(rc == 0 and lines and all(ln.count('"M01.2.s3"') <= 1 for ln in lines) and any("M01.2.s3" in ln for ln in lines) and "M01.2.s9" not in out,
            "候補: 共通の資料の手順は 1 回だけ出し、ほかのフローで描いた項目は出さない（描いた箱が候補になる）", out[-400:])
        snap.write_text(keep, encoding="utf-8")
        rc, out = run("sc_match.py", "material", "-w", d, "--row", "B03")
        say(rc == 0 and "F03-1:D-001" in out and "振込エラーがあるか" in out, "材料: 開いた下位フローの箱が、行の材料に入る", out[:300])
        rc, out = run("sc_axis.py", "plan", "-w", d)
        say(rc == 0 and "B03" in out, "材料: 行ごとの材料の大きさ", out[-300:])
    if wanted("番号"):
        edit(P_match(d, "B03"), split_task)
        recompute(d)
        nb = jload(d / "numbers.json")
        t = {x["id"]: x["no"] for x in jload(d / "rows" / "B03.result.json")["tasks"]}
        say(nb["assigned"].get("B03|t5") == "T-008" and t["t3"] == "T-006" and t["t4"] == "T-007",
            "番号: タスクを足しても、前からのタスクの番号は変わらない（足したタスクは次の番号）", f"{t}")
        shutil.copy2(base / "rows" / "B03.match.json", P_match(d, "B03"))
        shutil.copy2(base / "numbers.json", d / "numbers.json")
        recompute(d)
    if wanted("付け直し"):
        def renum(f):
            node(f, "P-003")["id"] = "P-004"
            for e in f["edges"]:
                e["from"], e["to"] = ["P-004" if x == "P-003" else x for x in (e["from"], e["to"])]
        edit(P_flow(d, "TK", "F03"), renum)
        edit(P_match(d, "B03"), lambda m: task(m, "t4")["content"][0]["evidence"].append({"site": "TK", "item": "M01.2.s4"}))   # 外した項目の根拠は、付け直さない
        edit(P_match(d, "B03"), lambda m: task(m, "t4")["absent"]["OS"].update(not_here=["TK/F03:P-003"]))                  # ほかの拠点の記載なしの、受け取る側の根拠
        rc, out = run("sc_collect.py", "-w", d, "--site", "TK", "--no-gate")
        moved = jload(P_match(d, "B03"))
        recv = task(moved, "t4")["absent"]["OS"].pop("not_here")
        edit(P_match(d, "B03"), lambda m: task(m, "t4")["absent"]["OS"].pop("not_here"))
        m = json.dumps(moved, ensure_ascii=False)
        rc2, out2 = check(d, "rows")
        say(rc == 0 and "F03:P-004" in m and "F03:P-003" not in m and '"item": "M01.2.s4"' in m and recv == ["TK/F03:P-004"] and rc2 == 0,
            "付け直し: 拠点の箱の番号が振り直されても、出典の指紋でタスクの鍵を付け直す（外した項目の根拠はそのまま）", out[-300:] + tail(out2))
        shutil.rmtree(d / "asis" / "TK")
        shutil.copytree(base / "asis" / "TK", d / "asis" / "TK")
        for name in ("snap/TK.json", "rows/B03.match.json", "axis.json"):
            shutil.copy2(base / name, d / name)
    if wanted("回答"):
        import openpyxl
        p = d / "out" / "拠点比較.xlsx"
        wb = openpyxl.load_workbook(p)
        ws = wb["確認事項一覧"]
        head = [c.value for c in ws[1]]
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "Q-003")
        ws.cell(row, head.index("回答") + 1).value = "経理担当が手で会計システムに記帳している"
        ws.cell(row, head.index("回答者") + 1).value = "大阪支店 経理担当"
        wb.save(p)
        rc, out = run("sc_book.py", "-w", d)
        ws = openpyxl.load_workbook(p)["確認事項一覧"]
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "Q-003")
        say(rc == 0 and ws.cell(row, head.index("回答") + 1).value == "経理担当が手で会計システムに記帳している"
            and ws.cell(row, head.index("回答者") + 1).value == "大阪支店 経理担当",
            "回答: 確認事項一覧の回答・回答者は、書き出し直しても残る", out[-300:])
        rc, out = check(d, "outputs")
        say(rc == 0, "回答: 回答を入れたブックを書き出し直しても、納品物の点検が FAIL 0", tail(out))
    shutil.rmtree(d)


def status_and_gate(base, root):
    d = copy(base, root, "gate")
    if wanted("区分"):
        e = d / "k1"
        shutil.copytree(d, e, ignore=shutil.ignore_patterns("k1"))
        def out_of_scope(a):
            for x in next(r for r in a["rows"] if r["id"] == "B01")["sites"]:
                if x["site"] != "TK":
                    x.update(relation="out_of_scope", parts=[], reason="この拠点は申請を本部で受ける")
        J("axis.json", out_of_scope)(e)
        M("B01", lambda m: [t["members"].update({s: [] for s in t["members"] if s != "TK"}) or t.update(absent={}) for t in m["tasks"]])(e)
        recompute(e)
        r = jload(e / "rows" / "B01.result.json")
        say(all(t["class"] == "拠点固有" for t in r["tasks"]), "区分: ほかの拠点が対象外の行の、1 拠点だけの作業は拠点固有（共通にしない）",
            f"{[(t['name'], t['class']) for t in r['tasks']]}")
        run("sc_drawio.py", "-w", e)
        fig = (e / "out" / "拠点比較フロー.drawio").read_text(encoding="utf-8")
        say('id="lane-OS.?none-h" value="（対象外）"' in fig and "対象外（この拠点では行わない業務・作業）" in fig,
            "区分: 対象外の拠点のレーンの見出しは（対象外）で、凡例に対象外の箱がある", re.findall(r'id="lane-OS[^"]*" value="[^"]*"', fig)[:3])
        shutil.rmtree(e)
    if wanted("対象外"):
        from sc_rows import Model
        import sc_common
        import sc_match
        import sc_verify
        e = d / "x1"                                                 # ほかの拠点の担当のレーンの箱
        shutil.copytree(d, e, ignore=shutil.ignore_patterns("x1"))
        before = jload(e / "rows" / "B03.result.json")
        cl0 = sc_verify.claims(Model(e, jload(e / "sites.json")), "B03")
        J("terms.json", other_finance)(e)
        M("B03", lambda m: task(m, "t4").update(content=[]))(e)
        recompute(e)
        r = jload(e / "rows" / "B03.result.json")
        x = {k["id"]: k for k in r["tasks"]}
        req = next((q for q in r["requests"] if q["site"] == "SG"), {"coarse": [], "absent": []})
        score = lambda res: [res["sites"]["SG"]["complexity"][b]["score"] for b in ("aligned", "raw")]
        say(x["t3"]["marks"]["SG"] == "対象外" and x["t4"]["marks"]["SG"] == "対象外" and "資料では、東京支店の経理部（東京支店）が行う" in x["t3"]["out"].get("SG", "")
            and x["t3"]["no"] not in req["coarse"] + req["absent"] and x["t4"]["class"] == "拠点固有" and score(r) == score(before),
            "対象外: ほかの拠点の担当のレーンの箱だけのタスクは対象外（区分・依頼に数えない。複雑性は同じ）",
            f"{[(k['id'], k['marks'], k['class'], k.get('out')) for k in r['tasks']]} {req}")
        say(x["t2"]["marks"]["SG"] == "○" and x["t2"]["class"] == "拠点固有", "対象外: どの拠点にも自拠点の箱が無いタスクは、ほかの拠点の箱を記載ありにする",
            f"{x['t2']['marks']} {x['t2']['class']}")
        mdl = Model(e, jload(e / "sites.json"))
        cl = sc_verify.claims(mdl, "B03")
        sheet = sc_verify.sheet_text(mdl, "B03", "V-9")
        say(cl["t3/SG"][1].get("owner") and cl["t3/SG"][0] != cl0["t3/SG"][0] and "owner" not in cl["t3/TK"][1] and cl["t3/TK"][0] == cl0["t3/TK"][0]
            and "どの箱も、東京支店の担当が行う" in sheet and "シンガポール拠点の手順書の Finance は、東京支店の経理部" in sheet,
            "対象外: ほかの拠点の担当の箱を含む主張は、担当をその拠点の人とした理由つきで照合に回す（ほかの主張の印は変わらない）", sheet[-700:])
        run("sc_drawio.py", "-w", e)
        run("sc_book.py", "-w", e)
        fig = (e / "out" / "拠点比較フロー.drawio").read_text(encoding="utf-8")
        say("対象外（東京支店の経理部（東京支店）が行う）" in fig and "ほかの拠点が行う作業" in fig and "dashPattern=6 3" in fig,
            "対象外: 図は、ほかの拠点の担当のレーンの箱を灰色にし、凡例に出す")
        import openpyxl
        ws = openpyxl.load_workbook(e / "out" / "拠点比較.xlsx")["拠点タスク一覧"]
        head = [c.value for c in ws[1]]
        row = next(rr for rr in ws.iter_rows(min_row=2, values_only=True) if rr[0] == x["t3"]["no"])
        say("対象外" not in str(row[head.index("差の種類")]) and "対象外: シンガポール拠点" in str(row[head.index("差の内容")])
            and str(row[head.index("シンガポール拠点")]).startswith("対象外"),
            "対象外: 表は、記号を対象外にし、わけを差の内容に出す（差の種類には出さない）", str(row))
        tt = openpyxl.load_workbook(e / "out" / "拠点比較.xlsx")["拠点間の用語対応"]
        why = next((rr for rr in tt.iter_rows(min_row=2, values_only=True) if "東京支店の経理部" in rr), ())
        say(any(str(v).startswith("東京支店の担当（ほかの拠点）。シンガポール拠点の手順書") for v in why),
            "対象外: 用語対応のシートの根拠の欄に、ほかの拠点の担当と出す", str(why))
        shutil.rmtree(e)
        e = d / "x2"                                                 # 自拠点の箱とほかの拠点の箱が混ざる
        shutil.copytree(d, e, ignore=shutil.ignore_patterns("x2"))
        had = [k["kind"] for k in jload(e / "rows" / "B02.result.json")["tasks"][0]["diffs"]]
        term("R05", lambda g: g.update(site="TK", reason="試験: 支店長を東京支店の人とする"))(e)
        recompute(e)
        now = [k["kind"] for k in jload(e / "rows" / "B02.result.json")["tasks"][0]["diffs"]]
        say("担当" in had and "担当" not in now, "対象外: 自拠点の箱もあるタスクは、差を自拠点の箱だけで見る（ほかの拠点の担当の差を出さない）", f"{had} → {now}")
        shutil.rmtree(e)
        e = d / "x3"                                                 # 記載なしに、この拠点の手順ではない根拠（not_here）
        shutil.copytree(d, e, ignore=shutil.ignore_patterns("x3"))
        M("B02", split_b02)(e)
        mdl = Model(e, jload(e / "sites.json"))
        t2 = next(k for k in mdl.tasks("B02") if k["id"] == "t2")
        ck = {s: [c[0] for c in sc_match.candidates(mdl, "B02", t2, s)] for s in ("OS", "SG")}
        def fill(m):
            a = task(m, "t2")["absent"]
            a["OS"].update(checked=sorted(set(ck["OS"]) | {"M01.2.s3"}), not_here=["M01.2.s3"], note="共通の資料の書き分けで、大阪支店では課長が承認しない")
            a["SG"].update(checked=ck["SG"])
        M("B02", fill)(e)
        recompute(e)
        rc, out = check(e, "rows")
        r = jload(e / "rows" / "B02.result.json")
        x = next(k for k in r["tasks"] if k["id"] == "t2")
        req = next((q for q in r["requests"] if q["site"] == "OS"), {"absent": []})
        say(rc == 0 and x["marks"]["OS"] == "対象外" and x["marks"]["SG"] == "—" and x["class"] == "拠点固有" and x["no"] not in req["absent"]
            and "共通の資料の書き分けで、東京支店・シンガポール拠点だけの手順" in x["out"]["OS"],
            "対象外: 記載なしに not_here を書くと対象外（依頼に出ない）", tail(out) + f" {x['marks']} {x.get('out')} {req}")
        mdl = Model(e, jload(e / "sites.json"))
        cl = sc_verify.claims(mdl, "B02")
        sheet = sc_verify.sheet_text(mdl, "B02", "V-9")
        quote = next(q["quote"] for q in jload(e / "snap" / "OS.json")["excluded"] if q["key"] == "M01.2.s3")
        say("not_here" in cl["t2/OS:absent"][1] and "not_here" not in cl["t2/SG:absent"][1] and "の作業は、この拠点の手順ではない" in sheet and quote and quote in sheet,
            "対象外: 照合シートに、この拠点の手順ではない主張と根拠の引用が出る（not_here の無い主張の印は変わらない）", sheet[-600:])
        shutil.rmtree(e)
        e = d / "x4"                                                 # 内容の差の根拠に、外した項目
        shutil.copytree(d, e, ignore=shutil.ignore_patterns("x4"))
        M("B01", lambda m: task(m, "t2")["content"][0]["evidence"].append({"site": "OS", "item": "M01.2.s3"}))(e)
        rc, out = check(e, "rows")
        sheet = sc_verify.sheet_text(Model(e, jload(e / "sites.json")), "B01", "V-9")
        say(rc == 0 and "大阪支店の根拠（フローに描かなかった、資料の項目）" in sheet and quote in sheet,
            "対象外: 内容の差の根拠に、拠点がフローから外した項目を書け、照合シートに引用が出る", tail(out))
        shutil.rmtree(e)
        e = d / "x5"                                                 # 受け取る側（ほかの拠点の箱が「シンガポール拠点へ送付する」）
        shutil.copytree(d, e, ignore=shutil.ignore_patterns("x5"))
        J("snap/TK.json", to_sg)(e)
        J("snap/OS.json", to_sg)(e)
        mdl = Model(e, jload(e / "sites.json"))
        rec = mdl.receiver(next(k for k in mdl.tasks("B01") if k["id"] == "t1"), "SG")
        M("B01", lambda m: task(m, "t1")["absent"]["SG"].update(not_here=rec))(e)
        recompute(e)
        rc, out = check(e, "rows")
        r = jload(e / "rows" / "B01.result.json")
        x = next(k for k in r["tasks"] if k["id"] == "t1")
        req = next((q for q in r["requests"] if q["site"] == "SG"), {"absent": []})
        sheet = sc_verify.sheet_text(Model(e, jload(e / "sites.json")), "B01", "V-9")
        say(rc == 0 and rec == ["TK/F01:P-001", "OS/F01:P-001"] and x["marks"]["SG"] == "対象外" and "（シンガポール拠点は受け取る側）" in x["out"]["SG"]
            and x["no"] not in req["absent"] and x["class"] == "共通" and "この拠点は受け取る側" in sheet and "東京支店の箱:" in sheet,
            "対象外: ほかの拠点の箱がこの拠点へ行う作業と書けば、受け取る側として対象外（区分は残りの拠点で決まり、照合シートに根拠の箱が出る）",
            tail(out) + f" {rec} {x['marks']} {x['class']} {x.get('out')}")
        sites = jload(e / "sites.json")
        say(sc_common.addressed("申請書を東京の経理部へ報告する", "TK", sites) and sc_common.addressed("申請書を東京支店に提出する", "TK", sites)
            and not sc_common.addressed("申請書を東京支店に保管する", "TK", sites) and not sc_common.addressed("東京支店の口座から支払う", "TK", sites),
            "対象外: 受け取る側は「へ」か「に＋受け渡しの動詞」で見る（「東京支店に保管する」は拾わない）")
        shutil.rmtree(e)
    if wanted("図の形"):
        import sc_drawio
        fig = (d / "out" / "拠点比較フロー.drawio").read_text(encoding="utf-8")
        lay = {p["id"]: p for p in jload(d / "out" / "拠点比較フロー.layout.json")["pages"]}
        n = lay["row-B01"]["nodes"]
        cx = lambda k: n[k][0] + n[k][2] / 2
        w = re.search(r'id="lane-TASK-h"[^>]*>\s*<mxGeometry x="[^"]*" y="[^"]*" width="([0-9.]+)"', fig)
        say(abs(cx("TK.F01.start") - cx("TK.F01.P-001")) < 1 and w and float(w.group(1)) == 190,
            "図の形: 開始・終了の丸は箱の真ん中にそろい、タスクの列はタスクの名前の幅", f"{cx('TK.F01.start')} {cx('TK.F01.P-001')} {w and w.group(1)}")
        b2 = lay["row-B02"]["nodes"]
        ends = [k for k in b2 if k.startswith("TK.") and ".end" in k]
        say(len(ends) == 2 and all(b2[k][1] > max(v[1] for kk, v in b2.items() if kk.startswith("TK.") and ".end" not in kk) for k in ends)
            and b2["TK.F02.end"][1] < b2["TK.F02.end2"][1],
            "図の形: 途中の終了も、いちばん下の帯に（ふつうの終了が上）", f"{[(k, b2[k][1]) for k in ends]}")
        e = d / "f1"
        shutil.copytree(d, e, ignore=shutil.ignore_patterns("f1"))
        J("snap/TK.json", lambda s: s["nodes"]["F02:end2"].update(label="終了"))(e)
        run("sc_drawio.py", "-w", e)
        b2 = {p["id"]: p for p in jload(e / "out" / "拠点比較フロー.layout.json")["pages"]}["row-B02"]
        tk = [k for k in b2["nodes"] if k.startswith("TK.") and ".end" in k]
        into = [x for x in b2["edges"] if x["to"] == "TK.F02.end"]
        say(tk == ["TK.F02.end"] and len(into) == 2, "図の形: 名前が同じ終了は、拠点の塊で 1 つにまとめる", f"{tk} {[(x['from'], x['to']) for x in into]}")
        shutil.rmtree(e)
        rows = sc_drawio.balanced_lines("T-001: 経費申請書を作成する", 138, 11)
        say(rows is not None and len(rows) == 2 and BD_units(rows[-1]) >= 3 and "".join(rows).replace(" ", "") == "T-001:経費申請書を作成する"
            and sc_drawio.balanced_lines("T-001: 申請を承認する", 138, 11) is None,
            "図の形: 最後の行が 1〜2 字になる名前だけ、機械が行に分ける", f"{rows}")
        mdl = sc_drawio.Model(d, jload(d / "sites.json"))
        say(mdl.lane_label("SG", "R01") == "申請者（Employee）" and mdl.lane_label("TK", "R01") == "申請者" and mdl.lane_label("OS", "R03") == "経理（経理担当）",
            "図の形: レーンの名前は、比べる用語の名前に拠点の言葉を添える", f"{mdl.lane_label('SG', 'R01')} {mdl.lane_label('OS', 'R03')}")
    if wanted("書体"):
        import zipfile
        fig = (d / "out" / "拠点比較フロー.drawio").read_text(encoding="utf-8")
        cells = [c for c in re.findall(r"<mxCell [^>]*>", fig) if 'vertex="1"' in c or 'edge="1"' in c]
        bad = [c for c in cells if "fontFamily=Yu Gothic UI;" not in c]

        def faces(path):
            xml = zipfile.ZipFile(path).read("xl/styles.xml").decode("utf-8")
            fonts = re.findall(r"<font>(.*?)</font>|<font/>", re.search(r"<fonts[^>]*>(.*?)</fonts>", xml, re.S).group(1), re.S)
            return [(re.search(r'<name val="([^"]*)"', f) or [None, None])[1] for f in fonts]
        book, tpl = faces(d / "out" / "拠点比較.xlsx"), faces(SKILL / "assets" / "拠点比較_テンプレート.xlsx")
        say(cells and not bad and set(book) == {"Yu Gothic UI"} and set(tpl) == {"Yu Gothic UI"},
            "書体: 図のどのセルも、表のブックとテンプレートの書体も Yu Gothic UI", f"{len(cells)} {bad[:2]} {set(book)} {set(tpl)}")
    if wanted("出典"):
        e = d / "o1"
        shutil.copytree(d, e, ignore=shutil.ignore_patterns("o1"))
        J("snap/SG.json", lambda sp: next(iter(sp["nodes"].values()))["src"].append("SG_明細表「支払」シート B12"))(e)
        run("sc_book.py", "-w", e)
        run("sc_drawio.py", "-w", e)
        rc, out = check(e, "outputs")
        say("行の ID" not in out and "資料 ID" not in out, "出典: 表の資料の出典のセル番地（シート B12）は、行の ID・資料 ID として止めない", tail(out))
        shutil.rmtree(e)
    if wanted("照合"):
        e = d / "v1"
        shutil.copytree(d, e, ignore=shutil.ignore_patterns("v1"))
        verify_report("B03", lambda r: r["findings"].append({"kind": "内容", "target": "t4/SG", "text": "記帳の時機の差も内容の差に足す"}))(e)
        J("review/B03.review.json", lambda r: r["findings"][0].update(decision="採用", action="内容の差を足した"))(e)
        M("B03", lambda m: task(m, "t4")["content"].append({"sites": ["SG"], "kind": "値", "text": "シンガポール拠点は月末にまとめて記帳する",
                                                            "evidence": [{"site": "TK", "node": "F03:S-001"}, {"site": "SG", "node": "F02:P-003"}]}))(e)
        recompute(e)
        rc, out = check(e, "verify")
        say("まだ反映していない" not in out and "判定・指摘が、照合の報告と違う" not in out,
            "照合: 採用した指摘を、同じタスクに内容の差を足して直せば、反映したとみなす", tail(out))
        shutil.rmtree(e)
    if wanted("進み具合"):
        rc, out = run("sc_status.py", "-w", d, "--fast")
        say(rc == 0 and "sc_deliver.py" in out and "未" in out.split("9  納品の門")[-1][:20], "進み具合: 書き出しまで済み、次は納品の門", out[-400:])
        P_match(d, "B02").rename(d / "B02.keep")
        rc, out = run("sc_status.py", "-w", d, "--fast")
        say("sc_match.py draft" in out and "--row B02" in out, "進み具合: タスクの対応の無い行は、次に下書きのコマンドを出す", out[-400:])
        (d / "B02.keep").replace(P_match(d, "B02"))
        rc, out = run("sc_run.py", "-w", d, "--dry-run")
        say(rc == 0 and "納品の門" in out, "進行役: 機械の工程を回し、納品の門で止まる", out[-300:])
    if wanted("配置"):
        e = d / "e1"
        shutil.copytree(d, e, ignore=shutil.ignore_patterns("e1"))
        (e / "out" / "拠点比較フロー.drawio").unlink()
        rc, out = run("sc_layout.py", "-w", e, "--no-png")
        say(rc == 1 and "が無い" in out, "配置: 図が無ければ止まる", tail(out))
        shutil.rmtree(e)
        e = d / "e2"
        shutil.copytree(d, e, ignore=shutil.ignore_patterns("e1", "e2"))
        overlap(e)
        rc, out = run("sc_layout.py", "-w", e, "--no-png")
        say(rc == 1 and "［経費申請］" in out, "配置: 箱が重なれば FAIL", tail(out))
        shutil.rmtree(e)
        rc, out = run("sc_layout.py", "-w", d, env=fake_drawio(root, ok=False))
        say(rc == 1 and "PNG に書き出せない" in out, "配置: PNG に書き出せなければ FAIL", tail(out))
        rc, out = run("sc_layout.py", "-w", d, "--no-png")
        say(rc == 0 and "0 FAIL / 0 WARN" in out, "配置: 架空の 3 拠点の図は、近すぎて並ぶ線も無い", tail(out))
        e = d / "e3"
        shutil.copytree(d, e, ignore=shutil.ignore_patterns("e1", "e2", "e3"))
        def close(lay):
            pg = next(p for p in lay["pages"] if p["id"] == "row-B01")
            pg["edges"] += [{"from": "a1", "to": "b1", "points": [[5000, 100], [5000, 200]]}, {"from": "a2", "to": "b2", "points": [[5006, 120], [5006, 220]]}]
        edit(e / "out" / "拠点比較フロー.layout.json", close)
        rc, out = run("sc_layout.py", "-w", e, "--no-png")
        say(rc == 0 and "6px の近さで 80px 並んで走る" in out, "配置: 近すぎて並ぶ線は WARN", tail(out))
        shutil.rmtree(e)
    if wanted("納品"):
        env = fake_drawio(root)
        dest = root / "dest"
        shutil.rmtree(dest, ignore_errors=True)
        rc, out = run("sc_deliver.py", "-w", d, "--to", dest, env=env)
        goods, log = dest / "納品物", dest / "作業の記録"
        pages = len(jload(d / "out" / "拠点比較フロー.layout.json")["pages"])
        ok = rc == 0 and (goods / "拠点比較.xlsx").exists() and (goods / "拠点比較フロー.drawio").exists() and len(list((goods / "png").glob("*.png"))) == pages \
            and (log / "比較の数.md").exists() and (log / "対応の照合_採否表.md").exists() and len(list(log.glob("対応の照合シート_*.md"))) == 3
        say(ok, "納品の門: 全部通れば、納品物と作業の記録を写す（draw.io の代わりの試験の道具で）", out[-500:])
        n = (log / "比較の数.md").read_text(encoding="utf-8") if (log / "比較の数.md").exists() else ""
        say("共通" in n and "拠点固有" in n and "支持率" in n, "納品の門: 比較の数（区分ごとの数・支持率）", n[:300])
        (goods / "メモ.txt").write_text("x", encoding="utf-8")
        rc, out = run("sc_deliver.py", "-w", d, "--to", dest, env=env)
        say(rc != 0 and "この門が写したもの以外のファイルがある" in out, "納品の門: 出力先にほかのファイルがあれば、上書きせずに止まる", out[-300:])
        edit(P_match(d, "B01"), lambda m: task(m, "t2").update(name="領収書を添えて申請書を出す"))
        rc, out = run("sc_deliver.py", "-w", d, "--check", env=env)
        say(rc != 0 and "照合したあとに変わった主張" in out, "納品の門: 照合のあとに変えたタスクがあれば止まる（止めた FAIL を見せる）", out[-400:])
        shutil.rmtree(dest, ignore_errors=True)
    shutil.rmtree(d)


def main():
    global ONLY
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-w", "--work", help="組み立てる場所（省くと一時フォルダ。終わったら消す）")
    ap.add_argument("--only", help="名前にこの語を含む試験だけ")
    a = ap.parse_args()
    ONLY = a.only
    root = pathlib.Path(a.work).expanduser().resolve() if a.work else pathlib.Path(tempfile.mkdtemp(prefix="sc_selftest_"))
    root.mkdir(parents=True, exist_ok=True)
    base = build_base(root)
    print(f"架空の 3 拠点 → {base}")
    if wanted("正しい") or wanted("schemas") or wanted("支持率 100"):
        good(base)
    if wanted("計算") or wanted("番号:"):
        expected(base)
    if wanted("部品"):
        unit()
    if wanted("揺れ"):
        stable(base, root)
    if wanted("表記"):
        wording(base)
    if wanted("基準なし"):
        no_baseline(base, root)
    cases(base, root)
    parts(base, root)
    status_and_gate(base, root)
    total = len([1 for _ in NG])
    print(f"\nNG {total}" + (f"（{'・'.join(NG[:8])}{' ほか' if total > 8 else ''}）" if total else ""))
    if not a.work:
        shutil.rmtree(root, ignore_errors=True)
    sys.exit(1 if NG else 0)


if __name__ == "__main__":
    main()
